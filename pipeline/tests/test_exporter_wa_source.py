from __future__ import annotations

from pipeline.agents.exporter import _fill_quantity


def test_fill_quantity_prefers_mapper_wa():
    """시트4 Zone = mapper_wa 우선, 없으면 parser zone fallback."""
    import openpyxl

    ws = openpyxl.Workbook().active
    rows = [
        {"mapper_wa": "FAB", "zone": None, "trade": "건축", "unit": "m2"},
        {"mapper_wa": None, "zone": "CUB", "trade": "토목", "unit": "m"},
        {"mapper_wa": "", "zone": "FAB", "trade": "건축", "unit": "m2"},
    ]
    _fill_quantity(ws, rows, __import__("datetime").date(2025, 9, 24))
    from pipeline.agents.exporter import COL_QUANTITY, DATA_START_ROW

    zcol = COL_QUANTITY["zone"]
    assert ws.cell(DATA_START_ROW + 0, zcol).value == "FAB"  # mapper_wa
    assert ws.cell(DATA_START_ROW + 1, zcol).value == "CUB"  # fallback parser zone
    assert ws.cell(DATA_START_ROW + 2, zcol).value == "FAB"  # 빈 mapper_wa -> parser zone
