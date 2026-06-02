"""Recipe 기반 추출 — 4개 시트 추출기.

parser.py 의 _parse_sheetN_* 4종을 그대로 포팅하되, 컬럼 위치를 _resolve_* 재호출이
아니라 recipe(columns/timeshift_blocks/final_flag_col) 로 주입. BP 코드 판정만
is_bp_code_v2 (Task 3) 로 교체. row loop·group 경계·company 탐색·skip·시간대 블록은
원본을 line-by-line 전사 (재설계 금지 — IRON LAW).

  - extract_sheet1: 시트1 출력인원 현황 -> BPSummary[] (col30 merge 경계)
  - extract_sheet2: 시트2 직종별 인원현황 -> JobAllocation[] (Mm01 앵커 경계)
  - extract_sheet3: 시트3 장비현황 -> EquipmentAllocation[] (equip_triple 경계)
  - extract_sheet4: 시트4 물량 -> Activity[] (bp_company_pair 경계)
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

from pipeline.agents.parser import (
    _build_activity,
    _disc_abbrev_to_kr,
    _extract_meta,
    _is_admin_job,
    _norm_header,
    _safe_int,
    _shift_dict_row,
    _shift_dict_sum,
)
from pipeline.agents.parser_post import (
    extract_zone,
    normalize_company,
    split_company_and_zone,
)
from pipeline.agents.recipe.bpcode import is_bp_code_v2
from pipeline.agents.recipe.model import FormatRecipe, SheetRecipe
from pipeline.agents.recipe.recognize import recognize_format
from pipeline.state import (
    Activity,
    BPSummary,
    EquipmentAllocation,
    JobAllocation,
    ParsedDailyReport,
)


def extract_sheet1(ws, sheet_recipe: SheetRecipe) -> list[BPSummary]:
    """시트1 (출력인원 현황) -> BPSummary[].

    _parse_sheet1_bp_summaries 전사:
      - 그룹 경계는 col30(=cols["today_incr_total"]) 컬럼의 merged range + solo row 흡수
      - 컬럼 위치는 recipe(sheet_recipe.columns/timeshift_blocks) 로 주입
      - BP 코드 판정만 is_bp_code_v2 로 교체, 그 외 row loop·skip·시간대 4블록은 line-by-line 전사
    """
    cols = sheet_recipe.columns
    blocks = sheet_recipe.timeshift_blocks
    tot_col = cols["today_incr_total"]
    cum_tot_col = cols.get("today_cum_total")
    bpc = cols["bp_company"]
    scol = cols.get("scope")

    # 1. col 30(=tot_col) merged 범위 수집 (값 있는 좌상단만)
    ad_merges: dict[int, int] = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_col == tot_col and mr.max_col == tot_col:
            top_val = ws.cell(mr.min_row, tot_col).value
            if top_val is None:
                continue
            ad_merges[mr.min_row] = mr.max_row

    # 2. solo row (merged 가 아니면서 tot_col 값 있는 row)
    solo_rows: list[int] = []
    for r in range(5, ws.max_row + 1):
        v = ws.cell(r, tot_col).value
        if v is None:
            continue
        if not any(s <= r <= e for s, e in ad_merges.items()):
            solo_rows.append(r)

    # 3. group range 결정 (solo+next 흡수)
    raw_groups: list[tuple[int, int]] = list(ad_merges.items())
    for r in solo_rows:
        next_r = r + 1
        absorbed = False
        if next_r <= ws.max_row:
            v5_next = ws.cell(next_r, bpc).value
            v30_next = ws.cell(next_r, tot_col).value
            in_other = any(s <= next_r <= e for s, e in ad_merges.items())
            if v5_next and v30_next is None and not in_other:
                raw_groups.append((r, next_r))
                absorbed = True
        if not absorbed:
            raw_groups.append((r, r))
    raw_groups.sort(key=lambda g: g[0])

    # 4. 각 group → BPSummary
    summaries: list[BPSummary] = []
    for start, end in raw_groups:
        bp_code: str | None = None
        company_raw: str | None = None
        col6_zone_raw: str | None = None
        work_text_today_parts: list[str] = []
        work_text_prev_parts: list[str] = []
        work_text_tomorrow_parts: list[str] = []

        for r in range(start, end + 1):
            v5 = ws.cell(r, bpc).value
            v6 = ws.cell(r, scol).value if scol else None
            v7 = ws.cell(r, cols["work_prev"]).value if "work_prev" in cols else None
            v8 = ws.cell(r, cols["work_today"]).value if "work_today" in cols else None
            v9 = ws.cell(r, cols["work_tomorrow"]).value if "work_tomorrow" in cols else None

            if is_bp_code_v2(v5):
                bp_code = str(v5).strip()
                if v6 is not None and col6_zone_raw is None:
                    col6_zone_raw = str(v6)
            elif v5:
                if company_raw is None:
                    company_raw = str(v5)
                if v6 is not None and col6_zone_raw is None:
                    col6_zone_raw = str(v6)

            if v7:
                work_text_prev_parts.append(str(v7).strip())
            if v8:
                work_text_today_parts.append(str(v8).strip())
            if v9:
                work_text_tomorrow_parts.append(str(v9).strip())

        company = normalize_company(split_company_and_zone(company_raw))
        zone = extract_zone(col6_zone_raw)

        today_incr_total = _safe_int(ws.cell(start, tot_col).value)
        today_cum_total = _safe_int(ws.cell(start, cum_tot_col).value) if cum_tot_col else 0

        # spec §2.2 skip 조건
        if bp_code is None and not company:
            continue  # sub-section 합계
        if not company and today_incr_total == 0 and today_cum_total == 0:
            continue  # 빈 placeholder

        workers_today_incr = _shift_dict_sum(ws, start, end, blocks["today_incr"])
        workers_today_cum = _shift_dict_sum(ws, start, end, blocks["today_cum"])
        workers_yesterday_incr = _shift_dict_sum(ws, start, end, blocks["yest_incr"])
        workers_yesterday_cum = _shift_dict_sum(ws, start, end, blocks["yest_cum"])

        summaries.append(
            {
                "bp_code": bp_code,
                "company": company or "",
                "disc_abbrev": None,
                "zone": zone,
                "scope": col6_zone_raw,
                "work_text_yesterday": " | ".join(work_text_prev_parts) or None,
                "work_text_today": " | ".join(work_text_today_parts) or None,
                "work_text_tomorrow": " | ".join(work_text_tomorrow_parts) or None,
                "workers_yesterday_incr": workers_yesterday_incr,
                "workers_yesterday_cum": workers_yesterday_cum,
                "workers_today_incr": workers_today_incr,
                "workers_today_cum": workers_today_cum,
                "today_incr_total": today_incr_total,
                "today_cum_total": today_cum_total,
            }
        )
    return summaries


def extract_sheet2(ws, sheet_recipe: SheetRecipe) -> list[JobAllocation]:
    """시트2 (직종별 인원현황) -> JobAllocation[].

    _parse_sheet2_job_allocations 전사:
      - has_final_flag 게이트: recipe 가 컬럼(final_flag_col)만 주입하고, 'O' 존재 여부는
        원본처럼 런타임 스캔으로 판정. 통합(col25 에 'O' 존재)은 group 시작·직종 row 둘 다
        게이트. 업체(final_flag_col=None)는 비활성 — 원본이 col25 무 'O' 일 때와 동일.
      - group 시작 = disc_code='Mm01' + trade·bp_company 채워있음 + 게이트 + is_bp_code_v2(bp)
      - 컬럼 위치·시간대 3블록은 recipe 로 주입, BP 코드 판정만 is_bp_code_v2 로 교체.
        그 외 group 경계·company 탐색·row loop·skip·관리자 filter 는 line-by-line 전사.
    """
    cols = sheet_recipe.columns
    blocks = sheet_recipe.timeshift_blocks
    ff = sheet_recipe.final_flag_col
    # 원본 _parse_sheet2_job_allocations 와 동일하게 런타임 'O' 스캔으로 게이트 활성 여부 결정.
    # recipe 는 컬럼(ff)만 주입 — ff 가 캐시/LLM 으로 설정됐어도 실제 'O' 가 없으면 게이트 비활성
    # (구조적 ff!=None 만으로 게이트하면 'O' 없는 변형 파일에서 전 row 조용히 drop → 오염).
    has_final_flag = ff is not None and any(ws.cell(r, ff).value == "O" for r in range(6, ws.max_row + 1))
    disc_code_col = cols["disc_code"]
    trade_col = cols["trade"]
    bp_col = cols["bp_company"]
    job_type_col = cols["job_type"]

    # 1. group 시작 row 식별
    starts: list[int] = []
    for r in range(6, ws.max_row + 1):
        c4 = ws.cell(r, disc_code_col).value
        if c4 != "Mm01":
            continue
        c2 = ws.cell(r, trade_col).value
        c3 = ws.cell(r, bp_col).value
        if c2 is None or c3 is None:
            continue
        if has_final_flag and ws.cell(r, ff).value != "O":
            continue
        if not is_bp_code_v2(c3):
            continue
        starts.append(r)

    groups: list[tuple[int, int]] = [
        (s, (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row) for i, s in enumerate(starts)
    ]

    jobs: list[JobAllocation] = []
    for start, end in groups:
        bp_code = str(ws.cell(start, bp_col).value).strip()
        discipline_raw = str(ws.cell(start, trade_col).value).strip()
        discipline = _disc_abbrev_to_kr(discipline_raw) or discipline_raw

        # 회사명: start 다음 5 row 안의 첫 bp_col 한글값 (BP 코드 아님)
        company_raw: str | None = None
        for r in range(start + 1, min(start + 6, end + 1)):
            v = ws.cell(r, bp_col).value
            if v is not None and not is_bp_code_v2(v):
                company_raw = str(v)
                break
        company = normalize_company(split_company_and_zone(company_raw)) or ""

        for r in range(start, end + 1):
            c4 = ws.cell(r, disc_code_col).value
            if c4 is None:
                continue
            if has_final_flag and ws.cell(r, ff).value != "O":
                continue
            job_type = str(ws.cell(r, job_type_col).value or "").strip()
            if not job_type:
                continue
            if _is_admin_job(job_type):
                continue
            disc_code = str(c4).strip()
            manager_code = disc_code if disc_code.startswith("Mm") else None
            worker_code = disc_code if not disc_code.startswith("Mm") else None
            workers_today_cum = _shift_dict_row(ws, r, blocks["today_cum"])
            jobs.append(
                {
                    "bp_code": bp_code,
                    "company": company,
                    "disc_abbrev": discipline,
                    "manager_code": manager_code,
                    "worker_code": worker_code,
                    "job_type": job_type,
                    "workers_yesterday_cum": _shift_dict_row(ws, r, blocks["yest_cum"]),
                    "workers_today_incr": _shift_dict_row(ws, r, blocks["today_incr"]),
                    "workers_today_cum": workers_today_cum,
                    "final_flag": "O",
                    "workers_cum_t": workers_today_cum.get("total", 0),
                }
            )

    return jobs


def extract_sheet3(ws, sheet_recipe: SheetRecipe) -> list[EquipmentAllocation]:
    """시트3 (장비현황) -> EquipmentAllocation[].

    _parse_sheet3_equipment_allocations 전사:
      - group 시작 = trade(col2)+bp_company(col3)+equip_big(col4) 동시 채워있음(=equip_triple
        앵커) + is_bp_code_v2(bp). 헤더 reprint row('장비 투입...') 회피.
      - has_final_flag 게이트: sheet2 와 동일하게 recipe 가 컬럼(final_flag_col)만 주입하고,
        'O' 존재 여부는 원본처럼 런타임 스캔으로 판정 (구조적 ff!=None 만으로 게이트하면
        'O' 없는 변형 파일에서 전 row 조용히 drop → 오염). 통합(col25 에 'O' 존재)은 장비 row
        에서 게이트, 업체(final_flag_col=None)는 비활성.
      - 컬럼 위치·시간대 3블록은 recipe 로 주입, BP 코드 판정만 is_bp_code_v2 로 교체.
        그 외 group 경계·company 탐색·row loop 은 line-by-line 전사. subtotal 없음.
    """
    cols = sheet_recipe.columns
    blocks = sheet_recipe.timeshift_blocks
    ff = sheet_recipe.final_flag_col
    # 원본 _parse_sheet3_equipment_allocations 와 동일하게 런타임 'O' 스캔으로 게이트 활성 여부 결정.
    has_final_flag = ff is not None and any(ws.cell(r, ff).value == "O" for r in range(6, ws.max_row + 1))
    trade_col = cols["trade"]
    bp_col = cols["bp_company"]
    equip_big_col = cols["equip_big"]
    equip_mid_col = cols["equip_mid"]
    equip_spec_col = cols["equip_spec"]

    starts: list[int] = []
    for r in range(6, ws.max_row + 1):
        c2 = ws.cell(r, trade_col).value
        c3 = ws.cell(r, bp_col).value
        c4 = ws.cell(r, equip_big_col).value
        if c2 is None or c3 is None or c4 is None:
            continue
        if not is_bp_code_v2(c3):
            continue
        # 헤더 reprint row 회피
        if "장비 투입" in str(c2):
            continue
        starts.append(r)

    groups: list[tuple[int, int]] = [
        (s, (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row) for i, s in enumerate(starts)
    ]

    eqs: list[EquipmentAllocation] = []
    for start, end in groups:
        bp_code = str(ws.cell(start, bp_col).value).strip()
        discipline_raw = str(ws.cell(start, trade_col).value).strip()
        discipline = _disc_abbrev_to_kr(discipline_raw) or discipline_raw

        company_raw: str | None = None
        for r in range(start + 1, min(start + 6, end + 1)):
            v = ws.cell(r, bp_col).value
            if v is not None and not is_bp_code_v2(v):
                company_raw = str(v)
                break
        company = normalize_company(split_company_and_zone(company_raw)) or ""

        for r in range(start, end + 1):
            c4 = ws.cell(r, equip_big_col).value
            if c4 is None:
                continue
            if has_final_flag and ws.cell(r, ff).value != "O":
                continue
            count_today_cum = _shift_dict_row(ws, r, blocks["today_cum"])
            eqs.append(
                {
                    "bp_code": bp_code,
                    "company": company,
                    "disc_abbrev": discipline,
                    "equipment_big": str(c4).strip(),
                    "equipment_mid": str(ws.cell(r, equip_mid_col).value or "").strip(),
                    "equipment_spec": (
                        str(ws.cell(r, equip_spec_col).value).strip()
                        if ws.cell(r, equip_spec_col).value is not None
                        else None
                    ),
                    "count_yesterday_cum": _shift_dict_row(ws, r, blocks["yest_cum"]),
                    "count_today_incr": _shift_dict_row(ws, r, blocks["today_incr"]),
                    "count_today_cum": count_today_cum,
                    "final_flag": "O",
                    "count_cum_t": count_today_cum.get("total", 0),
                }
            )
    return eqs


def extract_sheet4(ws, sheet_recipe: SheetRecipe) -> list[Activity]:
    """시트4 (물량) -> Activity[].

    _parse_sheet4_activities 전사:
      - 컬럼 위치는 recipe(sheet_recipe.columns) 로 주입 (통합/업체 흡수)
      - 통합포맷(company_col 존재): trade ffill + company 는 BP행에 종속(독립 ffill 아님)
      - 업체포맷(company_col 없음): BP 코드행~다음 코드행 group, company = group 내 첫 non-bp
      - Activity row = work_content 채워있는 row 만
      - zone = extract_zone(scope 명시 토큰만, 옵션 2 — scope ≠ wa)
    """
    cols = sheet_recipe.columns
    has_company_col = "company_col" in cols
    activities: list[Activity] = []
    if has_company_col:
        # 통합포맷: trade ffill + company 는 BP행에 종속(독립 ffill 아님)
        cur_trade = cur_bp = cur_company_raw = None
        for r in range(5, ws.max_row + 1):
            tv = ws.cell(r, cols["trade"]).value
            if tv is not None:
                cur_trade = str(tv)
            bv = ws.cell(r, cols["bp"]).value
            if bv is not None:
                cur_bp = str(bv).strip()
                cv = ws.cell(r, cols["company_col"]).value
                cur_company_raw = str(cv) if cv is not None else None
            wv = ws.cell(r, cols["work_content"]).value
            if wv is None:
                continue
            activities.append(_build_activity(ws, r, cols, cur_trade, cur_bp, cur_company_raw))
    else:
        # 업체포맷: BP사 컬럼 코드행~다음코드행 group, 회사명=group 내 첫 non-bp 값
        bpc = cols["bp"]
        code_rows = [r for r in range(5, ws.max_row + 1) if is_bp_code_v2(ws.cell(r, bpc).value)]
        for gi, start in enumerate(code_rows):
            end = (code_rows[gi + 1] - 1) if gi + 1 < len(code_rows) else ws.max_row
            bp_code = str(ws.cell(start, bpc).value).strip()
            company_raw = None
            for r in range(start, end + 1):
                v = ws.cell(r, bpc).value
                if v is not None and not is_bp_code_v2(v):
                    company_raw = str(v)
                    break
            cur_trade = None
            for r in range(5, start + 1):
                tv = ws.cell(r, cols["trade"]).value
                if tv is not None:
                    cur_trade = str(tv)
            for r in range(start, end + 1):
                if ws.cell(r, cols["work_content"]).value is None:
                    continue
                activities.append(_build_activity(ws, r, cols, cur_trade, bp_code, company_raw))
    return activities


# ---------------------------------------------------------------------------
# Task 10: 공개 API — extract_with_recipe + parse_with_recipe
# ---------------------------------------------------------------------------


def extract_with_recipe(wb, recipe: FormatRecipe, xlsx_path) -> ParsedDailyReport:
    """recipe 4시트 추출기 호출 + _extract_meta -> ParsedDailyReport.

    sheet_name 은 recipe.sheets[role].sheet_name 로 주입 (role: bp_summary/job/equip/activity).
    worker_outputs=[], report_meta=None 은 golden parse_daily_report 와 동일 (양식 2종 실측 일치).
    """
    bp_summaries = extract_sheet1(wb[recipe.sheets["bp_summary"].sheet_name], recipe.sheets["bp_summary"])
    job_allocations = extract_sheet2(wb[recipe.sheets["job"].sheet_name], recipe.sheets["job"])
    equipment_allocations = extract_sheet3(wb[recipe.sheets["equip"].sheet_name], recipe.sheets["equip"])
    activities = extract_sheet4(wb[recipe.sheets["activity"].sheet_name], recipe.sheets["activity"])
    report_date, project = _extract_meta(Path(xlsx_path), wb)
    return {
        "report_date": report_date,
        "project": project,
        "bp_summaries": bp_summaries,
        "job_allocations": job_allocations,
        "equipment_allocations": equipment_allocations,
        "activities": activities,
        "worker_outputs": [],
        "report_meta": None,
    }


def parse_with_recipe(xlsx_path) -> ParsedDailyReport:
    """xlsx 1개 -> recognize_format -> extract_with_recipe + 업체-전용 company mismatch review.

    업체(format_class=='업체') 파일에서만 파일명 토큰 vs 셀 company 를 대조해 mismatch 시
    해당 activity 의 requires_review=True (셀 company 는 SSOT, 덮어쓰기 X). 통합은 SKIP
    (파일명이 generic 이라 전체 review 처리되면 등가가 깨짐).
    """
    path = Path(xlsx_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        recipe = recognize_format(wb)
        report = extract_with_recipe(wb, recipe, path)
        if recipe.format_class == "업체":
            tokens = {_norm_header(t) for t in path.stem.split("_")}
            mismatched: set[str] = set()
            for a in report["activities"]:
                c = a["company"]
                norm_c = _norm_header(c)
                if not any(norm_c in tok for tok in tokens):
                    a["requires_review"] = True
                    mismatched.add(c)
            if mismatched:
                for company in sorted(mismatched):
                    print(
                        f"[recipe] company mismatch: 셀 '{company}' 가 파일명 '{path.name}' 에 없음 -> review",
                        file=sys.stderr,
                    )
    finally:
        wb.close()
    return report
