"""통합일보 09-25 → 회사별 건축/토목 업체파일 10개 생성 (덮어쓰기).

사용자 요청: 기존 10개 업체파일(전부 중원 UT배관 복사본)을 회사별 건축/토목 내용으로 덮어쓰기.
출처: docs/sample 통합작업일보_20250925 의 resolved 활동(회사별 건축/토목).
생성 형식: 통합포맷 시트4(모든 셀 명시 → ffill/병합 무관). sheet1/2/3 는 헤더만(워커는
delay/risk 에 무관 — forecaster agent.py:115 'delay_days qty 경로만').
생성 후 parse_daily_report 로 재검증.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path("/Users/kyungseoplee/Project Agent/CPIA")
sys.path.insert(0, str(ROOT))

from pipeline.agents.parser import parse_daily_report  # noqa: E402

SRC = ROOT / "docs/sample/통합작업일보_M15X Ph-3 IBL PJT 작업일보_20250925.xlsx"
OUT_DIR = ROOT / "data/demo_inputs/20250925"

# 파일명(회사, 공종표기) — 기존 10개 그대로. 회사명은 normalize 후 매칭.
TARGETS = [
    ("(유)성우이엔씨", "건축공사", "성우이엔씨"),
    ("㈜제일공사", "토공사", "제일공사"),
    ("㈜장평건설", "토목공사", "장평건설"),
    ("까뮤이앤씨", "외장공사", "까뮤이앤씨"),
    ("윤성이엔지", "외부공사", "윤성이엔지"),
    ("㈜엔아이스틸", "철골공사", "엔아이스틸"),
    ("㈜신성이엔지", "골조공사", "신성이엔지"),
    ("㈜에스비테크", "마감공사", "에스비테크"),
    ("㈜두제산업개발", "할석공사", "두제산업개발"),  # 통합일보에 없음 → 빈 파일
    ("남웅건설", "단열공사", "남웅건설"),
]

SHEET4 = "4.물량"
HEADER_ROW = 3
DATA_START = 5
# (col, header-text) — _resolve_sheet4_columns 가 텍스트로 매핑
COLS = {
    4: "공종",
    5: "BP사",
    6: "업체명",
    7: "Scope",
    8: "작업내용",
    9: "단위",
    10: "총물량",
    11: "전일누적",
    12: "금일실적",
    13: "금일누적",
    14: "F%",
}


def _build_wb(acts: list[dict]):
    wb = openpyxl.Workbook()
    # 시트1/2/3 — 헤더만(데이터 없음). 존재만 하면 parser 가 [] 로 처리.
    s1 = wb.active
    s1.title = "1.출력인원 현황"
    s1.cell(1, 2, "1. 출력인원 현황")
    s1.cell(3, 5, "BP사")
    s2 = wb.create_sheet("2.직종별 인원현황")
    s2.cell(1, 2, "2. 직종별 인원현황")
    s3 = wb.create_sheet("3.장비현황")
    s3.cell(1, 2, "3. 장비현황")
    # 시트4 — 헤더 + 활동행 (모든 셀 명시)
    s4 = wb.create_sheet(SHEET4)
    s4.cell(1, 2, "4. 물량")
    for col, txt in COLS.items():
        s4.cell(HEADER_ROW, col, txt)
    r = DATA_START
    for a in acts:
        s4.cell(r, 4, a.get("trade") or "")
        s4.cell(r, 5, a.get("bp_code") or "")
        s4.cell(r, 6, a.get("company") or "")
        s4.cell(r, 7, a.get("scope") or "")
        s4.cell(r, 8, a.get("work_content") or "")
        s4.cell(r, 9, a.get("unit") or "")
        s4.cell(r, 10, a.get("plan_total_qty"))
        s4.cell(r, 11, a.get("yesterday_cum_qty"))
        s4.cell(r, 12, a.get("today_qty"))
        s4.cell(r, 13, a.get("today_cum_qty"))
        s4.cell(r, 14, a.get("progress_percent"))
        r += 1
    return wb


def main() -> None:
    rep = parse_daily_report(str(SRC))
    acts = rep["activities"]
    by_co: dict[str, list[dict]] = defaultdict(list)
    for a in acts:
        if a.get("trade") in ("건축", "토목"):
            by_co[a.get("company") or ""].append(a)

    print(f"통합일보 건축/토목 활동: {sum(len(v) for v in by_co.values())} (회사 {len(by_co)})\n")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    # 기존 파일 제거
    for old in OUT_DIR.glob("*.xls*"):
        old.unlink()

    summary = []
    for raw_name, trade_label, norm_co in TARGETS:
        company_acts = by_co.get(norm_co, [])
        fname = f"업체작업일보_M15X_PH3_{raw_name}_{trade_label}_20250925.xlsx"
        out = OUT_DIR / fname
        wb = _build_wb(company_acts)
        wb.save(out)
        # 재검증: parse
        try:
            rp = parse_daily_report(str(out))
            n = len(rp["activities"])
            trades = sorted({x.get("trade") for x in rp["activities"]})
            cos = sorted({x.get("company") for x in rp["activities"]})
            ok = n == len(company_acts)
            summary.append((fname, len(company_acts), n, trades, cos, ok))
        except Exception as e:  # noqa: BLE001
            summary.append((fname, len(company_acts), f"PARSE FAIL: {e}", [], [], False))

    print("=== 생성 결과 (생성수 / 재파싱수 / trades / companies / OK) ===")
    allok = True
    for fname, src_n, parsed_n, trades, cos, ok in summary:
        allok = allok and ok
        print(f"  {'OK ' if ok else 'ERR'} {fname}")
        print(f"      src={src_n} parsed={parsed_n} trades={trades} companies={cos}")
    print(f"\n전체 OK: {allok}")


if __name__ == "__main__":
    main()
