"""9/25 데모 엑셀 재생성 — Path B: v1 이름·scope 유지 + 물량만 플랜 정합.

핵심(사용자 결정 2026-06-04):
  - 활동 식별자(work_content·scope·단위·장비) = 기존 v1 데모 그대로 (distinct 보존).
    v2(activity_master l7.5) 는 같은 이름 중복·scope없음(에폭시 ×21 동일 wbs_id) → 폐기.
  - 물량만 플랜 정합:
      A(총물량) = v1 값. 단 회사 ΣA > 플랜총량이면 플랜규모로 다운스케일(글로벌유명만 해당,
        v1 347k vs 플랜 40k → progress>1 로 CPI 붕괴 방지). scale=min(1, plan_A/Σv1A).
      B(전일누적) = A × 회사진행률 f × 활동별변주(±15%, clamp[0.30,0.95]).
        f = 시드 9/24 회사별 Σcum/Σplan (이력 연결·항상 매칭·l7.5 이름불일치 회피).
      C(금일실적) = (A−B) × c_frac(det 0.05~0.25) → D=B+C < A → 오버슈트0·잔여>0.
  - 인원 = 업체별 AC_c=EV_c/0.86 → 시트2 금일누계 주간 = AC_c/8 (CPI≈86%).
      EV_c=(ΣD/플랜A_full)×플랜MH_full.
  - 시트3 장비 = v1 그대로. 포맷 writer·템플릿 v2 재사용(파서 안전). OUT=v2/.
  - 재시드/DB write/마스터 변경 없음.

검증: 생성 후 parse_daily_report → 오버슈트0·잔여>0·이름·scope 보존 + CPI.
옵션: argv[1] = 단일 회사.
"""

from __future__ import annotations

import asyncio
import hashlib
import sys
import warnings
from datetime import date
from pathlib import Path

warnings.filterwarnings("ignore")
ROOT = Path("/Users/kyungseoplee/Project Agent/CPIA")
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from pipeline.scripts.gen_company_files_v2 import (  # noqa: E402
    TEMPLATE,
    _copy_style_row,
    _fill_sheet1,
    _fill_sheet2,
    _fill_sheet3,
    _set,
)
from pipeline.agents.parser import parse_daily_report  # noqa: E402
from pipeline.db import PGClient  # noqa: E402

PG = "postgresql://cpia:K7mP2xQ9nR4v@43.200.85.209:5532/cpia"
SITE = "m15x-ph3-ibl"
D24 = date(2025, 9, 24)
TARGET_CPI = 0.86
V1_DIR = ROOT / "data/demo_inputs/20250925"
OUT_DIR = ROOT / "data/demo_inputs/20250925/v2"
DATA0 = 5
STYLE_ROW = 7


def _det_frac(*parts: str) -> float:
    """결정성 [0,1) — 활동별 변주·c_frac 용."""
    h = int(hashlib.md5("|".join(parts).encode()).hexdigest(), 16)
    return (h % 100000) / 100000.0


def _fill_sheet4_pathb(ws, acts: list[dict], bp: str, company: str):
    """시트4 — v1 work_content/scope 유지 + Path B 물량. c4=공종 c5=BP/회사 c6=Scope
    c7=작업내용(v1) c8=단위 c9=A c10=B c11=C c12=D c13=F% c14=잔여 c15=G%."""
    for r in range(DATA0, ws.max_row + 1):
        for c in range(1, 20):
            cell = ws.cell(r, c)
            if cell.__class__.__name__ != "MergedCell":
                cell.value = None
    if not acts:
        return
    r = DATA0
    for i, a in enumerate(acts):
        _copy_style_row(ws, STYLE_ROW, r, 19)
        A = a["A"]; B = a["B"]; C = a["C"]; D = a["D"]
        _set(ws, r, 4, a.get("trade") or "")
        if i == 0:
            _set(ws, r, 5, bp)
        _set(ws, r, 6, a.get("scope") or "")
        _set(ws, r, 7, a.get("work_content") or "")
        _set(ws, r, 8, a.get("unit") or "")
        _set(ws, r, 9, round(A, 2) or None)
        _set(ws, r, 10, round(B, 2) or None)
        _set(ws, r, 11, round(C, 2))
        _set(ws, r, 12, round(D, 2) or None)
        _set(ws, r, 13, round(D / A, 4) if A else None)
        _set(ws, r, 14, round(A - D, 2) if A else None)
        _set(ws, r, 15, round(1 - D / A, 4) if A else None)
        r += 1
    # 회사명행 보장 (활동 1개여도 — parser.py:495-499)
    if len(acts) < 2:
        _copy_style_row(ws, STYLE_ROW, DATA0 + 1, 19)
    _set(ws, DATA0 + 1, 5, company)


def _calib_jobs(daytime_cum: int, daytime_incr: int, bp: str, company: str):
    z = {k: 0 for k in ("morning", "daytime", "overtime", "night", "all_night", "total")}
    tc = dict(z, daytime=daytime_cum, total=daytime_cum)
    ti = dict(z, daytime=daytime_incr, total=daytime_incr)
    yc = dict(z, daytime=max(0, daytime_cum - daytime_incr), total=max(0, daytime_cum - daytime_incr))
    job = {"company": company, "bp_code": bp, "final_flag": "O", "manager_code": None,
           "worker_code": "Wk01", "job_type": "기능공",
           "workers_today_cum": tc, "workers_today_incr": ti, "workers_yesterday_cum": yc}
    return job, tc, ti, yc


async def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    pg = PGClient(PG)
    await pg.connect()
    async with pg._pool.acquire() as con:
        # 회사 진행률 f (시드 9/24 Σcum/Σplan)
        fr = await con.fetch(
            """
            SELECT ma.company, sum(ma.today_cum_qty)::float cum, sum(ma.plan_total_qty)::float plan
            FROM mapper_activities ma JOIN pipeline_runs pr USING(run_id)
            WHERE pr.site_id=$1 AND ma.report_date=$2 GROUP BY ma.company
            """, SITE, D24,
        )
        company_f = {r["company"]: (r["cum"] / r["plan"]) for r in fr if r["plan"]}
        # 회사 플랜 합 (activity_master)
        pr2 = await con.fetch(
            "SELECT company, sum(plan_total_qty)::float a, sum(plan_mh_rev2)::float mh FROM activity_master WHERE site_id=$1 GROUP BY company", SITE
        )
        plan_A = {r["company"]: r["a"] for r in pr2}
        plan_MH = {r["company"]: r["mh"] for r in pr2}
    await pg.close()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    v1_files = sorted(V1_DIR.glob("*.xlsm"))
    if only is None:
        for old in OUT_DIR.glob("*.xls*"):
            old.unlink()

    print(f"=== Path B 재생성 (OUT={OUT_DIR}) ===")
    for v1path in v1_files:
        rep = parse_daily_report(str(v1path))
        v1acts = [a for a in rep["activities"] if a.get("plan_total_qty")]
        if not v1acts:
            continue
        company = v1acts[0].get("company") or ""
        if only and company != only:
            continue
        f = company_f.get(company, 0.7)
        pA = plan_A.get(company, 0) or 1
        sum_v1A = sum(float(a["plan_total_qty"]) for a in v1acts)
        scale = min(1.0, pA / sum_v1A) if sum_v1A else 1.0

        acts = []
        for a in v1acts:
            A = float(a["plan_total_qty"]) * scale
            wc = a.get("work_content") or ""
            sc = a.get("scope") or ""
            var = 0.85 + 0.30 * _det_frac(company, wc, sc)  # ±15%
            f_act = max(0.30, min(0.95, f * var))
            B = A * f_act
            c_frac = 0.05 + 0.20 * _det_frac(wc, sc, "c")
            C = (A - B) * c_frac
            D = B + C
            acts.append({"trade": a.get("trade") or "", "work_content": wc, "scope": sc,
                         "unit": a.get("unit") or "", "A": A, "B": B, "C": C, "D": D})

        sumD = sum(x["D"] for x in acts)
        sumC = sum(x["C"] for x in acts)
        ev_c = (sumD / pA) * plan_MH.get(company, 0)
        ac_c = ev_c / TARGET_CPI
        daytime_cum = max(1, round(ac_c / 8.0))
        daytime_incr = max(1, round((sumC / pA) * plan_MH.get(company, 0) / 8.0)) if pA else 1
        trade = acts[0]["trade"] or "건축"
        bp = "PH_X-001"

        wb = openpyxl.load_workbook(str(TEMPLATE), keep_vba=True)
        _fill_sheet4_pathb(wb["4.물량"], acts, bp, company)
        job, tc, ti, yc = _calib_jobs(daytime_cum, daytime_incr, bp, company)
        work_today = " | ".join(dict.fromkeys(x["work_content"] for x in acts if x["work_content"]))[:200]
        summary = {"company": company, "bp_code": bp, "work_text_today": work_today,
                   "workers_today_cum": tc, "workers_today_incr": ti, "workers_yesterday_cum": yc,
                   "today_incr_total": ti["total"], "today_cum_total": tc["total"]}
        _fill_sheet1(wb["1.출력인원 현황"], [summary], [job], company, trade)
        _fill_sheet2(wb["2.직종별 인원현황"], [job], company, trade)
        _fill_sheet3(wb["3.장비현황"], rep.get("equipment_allocations", []), company, trade)
        wb.save(str(OUT_DIR / v1path.name))
        print(f"  OK {company}: act={len(acts)} scale={scale:.3f} f={f:.2f} ΣA={sum(x['A'] for x in acts):.0f} "
              f"ΣD={sumD:.0f} EV_c={ev_c:.0f} 주간cum={daytime_cum} → {v1path.name}")

    print("\n생성 완료. 검증 별도.")


if __name__ == "__main__":
    asyncio.run(main())
