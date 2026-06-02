"""exporter _safe_set MergedCell-safe 쓰기 회귀 가드.

e2e 발견: render_xlsx → _fill_occupation 이 통합리포트 템플릿의 병합 셀에
ws.cell(r,c,value) 로 쓰다 'MergedCell' object attribute 'value' is read-only 크래시.
_safe_set 은 멤버 셀이면 병합 범위 anchor 로 우회 쓴다.
"""

from __future__ import annotations

import openpyxl

from pipeline.agents.exporter import _safe_set


def test_safe_set_redirects_merged_member_to_anchor() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B2:D2")  # B2=anchor, C2/D2 = MergedCell(read-only)
    # 멤버 셀 C2(row2,col3) 에 쓰기 — 크래시 없이 anchor B2(row2,col2) 로 들어가야 함.
    _safe_set(ws, 2, 3, "X")
    assert ws.cell(2, 2).value == "X"


def test_safe_set_writes_plain_cell_directly() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    _safe_set(ws, 5, 5, 42)
    assert ws.cell(5, 5).value == 42
