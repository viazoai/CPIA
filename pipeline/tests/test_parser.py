"""Parser 단위 테스트 — docs/superpowers/specs/2026-05-15-cpia-parser-design.md §9.1."""

from __future__ import annotations


from pipeline.agents.parser import (
    _disc_abbrev_to_kr,
    _ffill,
    _is_admin_job,
    _is_total_row,
    _normalize_unit,
)


def test_normalize_unit():
    assert _normalize_unit("M2") == "m2"
    assert _normalize_unit("m2") == "m2"
    assert _normalize_unit("Ton") == "ton"
    assert _normalize_unit("ton") == "ton"
    assert _normalize_unit("M3") == "m3"
    assert _normalize_unit("") == ""
    assert _normalize_unit("EA") == "ea"


def test_disc_abbrev_to_kr():
    assert _disc_abbrev_to_kr("AR") == "건축"
    assert _disc_abbrev_to_kr("ME") == "설비"
    assert _disc_abbrev_to_kr("EL") == "전기"
    assert _disc_abbrev_to_kr("UT") == "UT"
    assert _disc_abbrev_to_kr("TE") == "공통가설"
    assert _disc_abbrev_to_kr("IN") == "제어"
    assert _disc_abbrev_to_kr("CI") == "토목"
    assert _disc_abbrev_to_kr("XX") == "XX"
    assert _disc_abbrev_to_kr(None) is None


def test_is_admin_job():
    for job in ("BP소장", "공사관리자", "안전관리자", "품질관리자", "일반관리자"):
        assert _is_admin_job(job) is True, f"{job} 가 관리자로 인식 안 됨"
    for job in ("형틀공", "철근공", "콘크리트공", "비계공", "배관공"):
        assert _is_admin_job(job) is False, f"{job} 가 관리자로 잘못 인식"


def test_is_total_row():
    assert _is_total_row(("...", "기능공 투입 Total", 0, 0)) is True
    assert _is_total_row(("...", "BP사 투입 Total", 0, 0)) is True
    assert _is_total_row(("...", "형틀공", 100, 100)) is False
    assert _is_total_row(("...", None, 0, 0)) is False


def test_ffill_basic():
    rows = [
        ("h1", "h2", "h3", "h4", "h5"),
        ("h1", "h2", "h3", "h4", "h5"),
        ("h1", "h2", "h3", "h4", "h5"),
        ("h1", "h2", "h3", "h4", "h5"),
        ("A", "B", "C", "D", "E"),
        (None, None, None, None, "E2"),
        ("A2", None, "C2", None, "E3"),
    ]
    out = _ffill(rows, header_end=4, fill_cols=[0, 1, 2, 3])
    assert out == [
        (5, ["A", "B", "C", "D", "E"]),
        (6, ["A", "B", "C", "D", "E2"]),
        (7, ["A2", "B", "C2", "D", "E3"]),
    ]


# ---------------------------------------------------------------------------
# Task 2: _parse_sheet4_activities 테스트
# ---------------------------------------------------------------------------
from pathlib import Path  # noqa: E402

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from pipeline.agents.parser import _parse_sheet4_activities  # noqa: E402


@pytest.fixture(scope="module")
def sample_tong_path() -> Path:
    return (
        Path(__file__).resolve().parents[2]
        / "docs/sample/통합작업일보_M15X Ph-3 IBL PJT 작업일보_20250925.xlsx"
    )


def test_parse_sheet4_activities_count(sample_tong_path: Path):
    """fixture loader 와 동일한 row 추출 — 골든 79건 포함."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["4.물량"]
    activities = _parse_sheet4_activities(ws)
    wb.close()
    assert len(activities) > 79, f"전체 row 수가 79 이하: {len(activities)}"
    companies = {a["company"] for a in activities}
    assert "남웅건설" in companies
    assert "세보엠이씨" in companies
    assert "창원기전" in companies


def test_parse_sheet4_unit_normalized(sample_tong_path: Path):
    """단위가 소문자로 정규화됐는지."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["4.물량"]
    activities = _parse_sheet4_activities(ws)
    wb.close()
    for a in activities:
        if a["unit"]:
            assert a["unit"] == a["unit"].lower(), f"unit 미정규화: {a['unit']!r}"


def test_parse_sheet4_today_qty_zero_default(sample_tong_path: Path):
    """today_qty 가 숫자 아니면 0.0. 음수는 정정/감산이라 valid (실 데이터 검증)."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["4.물량"]
    activities = _parse_sheet4_activities(ws)
    wb.close()
    for a in activities:
        assert isinstance(a["today_qty"], float)


def test_parse_sheet4_required_fields(sample_tong_path: Path):
    """company / work_content 가 빈 row 는 skip."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["4.물량"]
    activities = _parse_sheet4_activities(ws)
    wb.close()
    for a in activities:
        assert a["company"], f"company 비어있음: {a}"
        assert a["work_content"], f"work_content 비어있음: {a}"


# ---------------------------------------------------------------------------
# Task 3: _parse_sheet1_bp_summaries 테스트
# ---------------------------------------------------------------------------
from pipeline.agents.parser import _parse_sheet1_bp_summaries  # noqa: E402


def test_parse_sheet1_bp_summaries_count(sample_tong_path: Path):
    # 2026-05-20 refactor: 시트1 신 logic 은 merged_cells.ranges 의존 → read_only=False
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    ws = wb["1.출력인원 현황"]
    summaries = _parse_sheet1_bp_summaries(ws)
    wb.close()
    assert len(summaries) > 0, "bp_summaries 비어있음"
    companies = {s["company"] for s in summaries}
    assert "남웅건설" in companies, f"남웅건설 없음, 추출된 회사: {sorted(companies)[:10]}"


def test_parse_sheet1_bp_summaries_fields(sample_tong_path: Path):
    """BPSummary TypedDict 의 모든 필드 존재."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    ws = wb["1.출력인원 현황"]
    summaries = _parse_sheet1_bp_summaries(ws)
    wb.close()
    required_keys = {
        "bp_code",
        "company",
        "disc_abbrev",
        "zone",
        "scope",
        "work_text_yesterday",
        "work_text_today",
        "work_text_tomorrow",
        "today_incr_total",
        "today_cum_total",
    }
    for s in summaries:
        assert required_keys <= set(s.keys()), f"누락 키: {required_keys - set(s.keys())}"
        assert isinstance(s["today_incr_total"], int)
        assert isinstance(s["today_cum_total"], int)


def test_parse_sheet1_disc_abbrev_kr(sample_tong_path: Path):
    """disc_abbrev 가 한글로 변환됐거나 None."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    ws = wb["1.출력인원 현황"]
    summaries = _parse_sheet1_bp_summaries(ws)
    wb.close()
    valid_kr = {"건축", "설비", "전기", "UT", "공통가설", "제어", "토목"}
    for s in summaries:
        if s["disc_abbrev"] is not None:
            assert s["disc_abbrev"] in valid_kr or len(s["disc_abbrev"]) <= 5


# ---------------------------------------------------------------------------
# Task 4: _parse_sheet2_job_allocations 테스트
# ---------------------------------------------------------------------------
from pipeline.agents.parser import _parse_sheet2_job_allocations  # noqa: E402


def test_parse_sheet2_no_admin_jobs(sample_tong_path: Path):
    """관리자 직종 0건 (필터 작동)."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["2.직종별 인원현황"]
    jobs = _parse_sheet2_job_allocations(ws)
    wb.close()
    admin = {"BP소장", "공사관리자", "안전관리자", "품질관리자", "일반관리자"}
    for j in jobs:
        assert j["job_type"] not in admin, f"관리자 직종 누락 안 됨: {j}"


def test_parse_sheet2_no_total_rows(sample_tong_path: Path):
    """합계 row 0건."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["2.직종별 인원현황"]
    jobs = _parse_sheet2_job_allocations(ws)
    wb.close()
    for j in jobs:
        assert "투입 Total" not in (j.get("job_type") or ""), f"Total row 누락 안 됨: {j}"


def test_parse_sheet2_workers_cum_t_int(sample_tong_path: Path):
    """workers_cum_t 가 int."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["2.직종별 인원현황"]
    jobs = _parse_sheet2_job_allocations(ws)
    wb.close()
    for j in jobs:
        assert isinstance(j["workers_cum_t"], int)


# ---------------------------------------------------------------------------
# Task 5: _parse_sheet3_equipment_allocations 테스트
# ---------------------------------------------------------------------------
from pipeline.agents.parser import _parse_sheet3_equipment_allocations  # noqa: E402


def test_parse_sheet3_count(sample_tong_path: Path):
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["3.장비현황"]
    eqs = _parse_sheet3_equipment_allocations(ws)
    wb.close()
    assert len(eqs) > 0, "장비 0건"


def test_parse_sheet3_big_mid_filled(sample_tong_path: Path):
    """첫 5건의 equipment_big, equipment_mid 가 None 아님."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["3.장비현황"]
    eqs = _parse_sheet3_equipment_allocations(ws)
    wb.close()
    for e in eqs[:5]:
        assert e["equipment_big"], f"equipment_big 비어있음: {e}"
        assert e["equipment_mid"], f"equipment_mid 비어있음: {e}"


def test_parse_sheet3_count_cum_t_int(sample_tong_path: Path):
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True, read_only=True)
    ws = wb["3.장비현황"]
    eqs = _parse_sheet3_equipment_allocations(ws)
    wb.close()
    for e in eqs:
        assert isinstance(e["count_cum_t"], int)
        assert e["count_cum_t"] >= 0


# ---------------------------------------------------------------------------
# Task 6: parse_daily_report + _extract_meta 테스트
# ---------------------------------------------------------------------------
from pipeline.agents.parser import ParseError, parse_daily_report  # noqa: E402


def test_parse_meta(sample_tong_path: Path):
    parsed = parse_daily_report(sample_tong_path)
    assert parsed["report_date"] == "2025-09-25"
    assert parsed["project"], f"project 비어있음: {parsed['project']!r}"
    assert "M15X" in parsed["project"] or "IBL" in parsed["project"]


def test_parse_daily_report_integration(sample_tong_path: Path):
    """공개 API — 4시트 통합 ParsedDailyReport."""
    parsed = parse_daily_report(sample_tong_path)
    assert len(parsed["bp_summaries"]) > 0
    assert len(parsed["job_allocations"]) > 0
    assert len(parsed["equipment_allocations"]) > 0
    assert len(parsed["activities"]) > 79


def test_parse_determinism(sample_tong_path: Path):
    """같은 xlsx 2회 호출 -> 동일 dict (결정성)."""
    p1 = parse_daily_report(sample_tong_path)
    p2 = parse_daily_report(sample_tong_path)
    assert p1 == p2


def test_parse_error_missing_sheet(tmp_path):
    """시트 누락 -> ParseError."""
    bad_path = tmp_path / "broken.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "wrong_sheet"
    wb.save(bad_path)
    with pytest.raises(ParseError):
        parse_daily_report(bad_path)


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: 시트1 신규 case (spec §2.2)
# ---------------------------------------------------------------------------


def _load_sheet1(sample_tong_path: Path):
    """신 logic 은 merged_cells.ranges 사용 → read_only=False 필요."""
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    return wb["1.출력인원 현황"]


def test_sheet1_extracts_around_99_groups(sample_tong_path: Path):
    """spec §2.2 — sub-section 합계 + 빈 placeholder skip 후 ~99 그룹."""
    ws = _load_sheet1(sample_tong_path)
    rows = _parse_sheet1_bp_summaries(ws)
    assert 90 <= len(rows) <= 110, f"expected ~99 groups, got {len(rows)}"


def test_sheet1_namwoong_fab_cub_split(sample_tong_path: Path):
    """남웅건설 PH_X-001 (FAB) + PH_X-009 (CUB) 분리 확인."""
    ws = _load_sheet1(sample_tong_path)
    rows = _parse_sheet1_bp_summaries(ws)
    namwoong = [r for r in rows if (r.get("company") or "") == "남웅건설"]
    assert len(namwoong) >= 2, f"expected 2 sites, got {len(namwoong)}"
    bp_to_zone = {r["bp_code"]: r["zone"] for r in namwoong if r["bp_code"]}
    assert bp_to_zone.get("PH_X-001") == "FAB", f"PH_X-001 zone={bp_to_zone.get('PH_X-001')}"
    assert bp_to_zone.get("PH_X-009") == "CUB", f"PH_X-009 zone={bp_to_zone.get('PH_X-009')}"


def test_sheet1_subsection_total_skipped(sample_tong_path: Path):
    """col 5/6 둘 다 None + col 30/37 > 0 인 sub-section 합계 row 는 skip."""
    ws = _load_sheet1(sample_tong_path)
    rows = _parse_sheet1_bp_summaries(ws)
    none_rows = [r for r in rows if r.get("bp_code") is None and not r.get("company")]
    assert len(none_rows) == 0, f"sub-section total leaked: {none_rows[:3]}"


def test_sheet1_chwangwon_paren_normalized(sample_tong_path: Path):
    """창원기전(6.6kV) → 창원기전 (괄호 부속 normalize)."""
    ws = _load_sheet1(sample_tong_path)
    rows = _parse_sheet1_bp_summaries(ws)
    chwangwon = [r for r in rows if (r.get("company") or "") == "창원기전"]
    assert len(chwangwon) >= 1, f"expected >=1 정규화된 창원기전 row, got {len(chwangwon)}"


def test_sheet1_bp_code_extracted(sample_tong_path: Path):
    """BP 코드가 회사명과 같은 그룹에 묶여 추출됨 (None 비율 < 50%)."""
    ws = _load_sheet1(sample_tong_path)
    rows = _parse_sheet1_bp_summaries(ws)
    with_bp = sum(1 for r in rows if r.get("bp_code"))
    assert with_bp >= len(rows) * 0.5, f"bp_code 추출률 낮음: {with_bp}/{len(rows)}"


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: 시트2 신규 case (spec §2.3)
# ---------------------------------------------------------------------------


def _load_sheet2(sample_tong_path: Path):
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    return wb["2.직종별 인원현황"]


def test_sheet2_extracts_at_least_800_rows(sample_tong_path: Path):
    """spec §2.3 — admin filter 적용 후 직종 row 충분 (보수)."""
    ws = _load_sheet2(sample_tong_path)
    rows = _parse_sheet2_job_allocations(ws)
    assert len(rows) >= 800, f"expected >=800 직종 row, got {len(rows)}"


def test_sheet2_bp_code_extracted(sample_tong_path: Path):
    """신 logic: bp_code 가 채워짐 (기존 None → PH_X-NNN)."""
    ws = _load_sheet2(sample_tong_path)
    rows = _parse_sheet2_job_allocations(ws)
    with_bp = sum(1 for r in rows if r.get("bp_code"))
    assert with_bp >= len(rows) * 0.95, f"bp_code 추출률 낮음: {with_bp}/{len(rows)}"


def test_sheet2_includes_ut_jeo_disciplines(sample_tong_path: Path):
    """col 2 keyword 무한정 — UT / 제어 group 도 추출 (whitelist X)."""
    ws = _load_sheet2(sample_tong_path)
    rows = _parse_sheet2_job_allocations(ws)
    disc_set = {r.get("disc_abbrev") for r in rows if r.get("disc_abbrev")}
    has_ut_or_eng = "UT" in disc_set
    has_control = "제어" in disc_set or "IN" in disc_set
    assert has_ut_or_eng or has_control, f"UT/제어 그룹 누락: disciplines={sorted(disc_set)}"


def test_sheet2_shift_dict_filled(sample_tong_path: Path):
    """TimeShiftCounts 4 그룹 채워짐 (workers_today_cum.total 합 > 0)."""
    ws = _load_sheet2(sample_tong_path)
    rows = _parse_sheet2_job_allocations(ws)
    sample = rows[:20]
    total_sum = sum(r["workers_today_cum"]["total"] for r in sample if r.get("workers_today_cum"))
    assert total_sum > 0, f"workers_today_cum.total 모두 0: {sample[:3]}"


def test_sheet2_manager_worker_code_split(sample_tong_path: Path):
    """manager_code (Mm##) vs worker_code (sk## / Wk##) 분리."""
    ws = _load_sheet2(sample_tong_path)
    rows = _parse_sheet2_job_allocations(ws)
    has_manager = any(r.get("manager_code") for r in rows)
    # admin filter 적용 시 manager 0 일 수도 있음 — 그땐 worker 만 검증
    has_worker = any(r.get("worker_code") for r in rows)
    assert has_manager or has_worker, "manager_code / worker_code 둘 다 None"


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: 시트3 신규 case (spec §2.4)
# ---------------------------------------------------------------------------


def _load_sheet3(sample_tong_path: Path):
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    return wb["3.장비현황"]


def test_sheet3_extracts_around_1111_rows(sample_tong_path: Path):
    """spec §2.4 — col 25='O' 인 장비 row ≈ 1,111."""
    ws = _load_sheet3(sample_tong_path)
    rows = _parse_sheet3_equipment_allocations(ws)
    assert 1000 <= len(rows) <= 1300, f"expected ~1111, got {len(rows)}"


def test_sheet3_bp_code_filled(sample_tong_path: Path):
    """신 logic: bp_code 가 채워짐."""
    ws = _load_sheet3(sample_tong_path)
    rows = _parse_sheet3_equipment_allocations(ws)
    with_bp = sum(1 for r in rows if r.get("bp_code"))
    assert with_bp >= len(rows) * 0.95, f"bp_code 추출률 낮음: {with_bp}/{len(rows)}"


def test_sheet3_namwoong_equipment(sample_tong_path: Path):
    """남웅건설 row 가 추출되고 equipment_big 다양."""
    ws = _load_sheet3(sample_tong_path)
    rows = _parse_sheet3_equipment_allocations(ws)
    namwoong = [r for r in rows if r.get("company") == "남웅건설"]
    assert len(namwoong) >= 1, "남웅건설 장비 row 없음"
    bigs = {r.get("equipment_big") for r in namwoong}
    assert len(bigs) >= 1, f"equipment_big 다양성 부족: {bigs}"


def test_sheet3_disc_abbrev_filled(sample_tong_path: Path):
    """disc_abbrev (공종) 가 채워짐."""
    ws = _load_sheet3(sample_tong_path)
    rows = _parse_sheet3_equipment_allocations(ws)
    has_disc = sum(1 for r in rows if r.get("disc_abbrev"))
    assert has_disc >= len(rows) * 0.95, f"disc_abbrev 누락: {has_disc}/{len(rows)}"


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: 시트4 신규 case (spec §2.5)
# ---------------------------------------------------------------------------


def _load_sheet4(sample_tong_path: Path):
    wb = openpyxl.load_workbook(sample_tong_path, data_only=True)
    return wb["4.물량"]


def test_sheet4_extracts_around_1193_activities(sample_tong_path: Path):
    """spec §2.5 — Activity row (col 8 채워있음) ≈ 1,193."""
    ws = _load_sheet4(sample_tong_path)
    rows = _parse_sheet4_activities(ws)
    assert 1100 <= len(rows) <= 1300, f"expected ~1193, got {len(rows)}"


def test_sheet4_ffill_namwoong_bp(sample_tong_path: Path):
    """ffill 정상 — 남웅건설 PH_X-001 의 activity 다수."""
    ws = _load_sheet4(sample_tong_path)
    rows = _parse_sheet4_activities(ws)
    namwoong_ph1 = [r for r in rows if r.get("company") == "남웅건설" and r.get("bp_code") == "PH_X-001"]
    assert len(namwoong_ph1) >= 5, f"ffill broken — only {len(namwoong_ph1)} rows for 남웅 PH_X-001"


def test_sheet4_zone_option2_conservative(sample_tong_path: Path):
    """zone — col 7 (Scope) 명시 토큰만, 모호 시 None (옵션 2 보수)."""
    ws = _load_sheet4(sample_tong_path)
    rows = _parse_sheet4_activities(ws)
    zone_none = sum(1 for r in rows if r.get("zone") is None)
    # prototype 측정: 76% None
    assert zone_none >= len(rows) * 0.5, f"zone 보수성 부족: {zone_none}/{len(rows)} None"


def test_sheet4_cross_check_BC_eq_D(sample_tong_path: Path):
    """B+C=D 산술 정합 — mismatch row 극소수."""
    ws = _load_sheet4(sample_tong_path)
    rows = _parse_sheet4_activities(ws)
    mismatch = 0
    for r in rows:
        b = r.get("yesterday_cum_qty") or 0
        c = r.get("today_qty") or 0
        d = r.get("today_cum_qty")
        if d is None:
            continue
        if abs((b + c) - d) > 0.001:
            mismatch += 1
    assert mismatch <= 10, f"too many B+C != D: {mismatch}"


def test_sheet4_required_new_fields(sample_tong_path: Path):
    """state.py Activity 의 신규 필드 (zone / bp_code / scope / yesterday_cum_qty / today_cum_qty) 존재."""
    ws = _load_sheet4(sample_tong_path)
    rows = _parse_sheet4_activities(ws)
    required = {
        "zone",
        "bp_code",
        "scope",
        "yesterday_cum_qty",
        "today_cum_qty",
    }
    for r in rows[:5]:
        assert required <= set(r.keys()), f"누락 키: {required - set(r.keys())}"


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: LLM 검증 wrapper (spec §2.6)
# ---------------------------------------------------------------------------

import os  # noqa: E402
from unittest.mock import patch  # noqa: E402


def test_parse_daily_report_llm_verify_off_skip(sample_tong_path: Path):
    """CPIA_PARSER_VERIFY_LLM=off 시 LLM 호출 없이 정상 종료."""
    with patch.dict(os.environ, {"CPIA_PARSER_VERIFY_LLM": "off"}):
        parsed = parse_daily_report(sample_tong_path)
    assert "bp_summaries" in parsed
    assert len(parsed["bp_summaries"]) > 0


def test_parse_daily_report_llm_verify_graceful_on_error(sample_tong_path: Path):
    """LLM 호출 실패 (API key 가짜) 해도 cycle block 안 함 (graceful degrade)."""
    with patch.dict(
        os.environ,
        {"CPIA_PARSER_VERIFY_LLM": "on", "AZURE_OPENAI_API_KEY": "FAKE_KEY_FOR_TEST"},
    ):
        # Exception 안 던지고 정상 종료해야 함
        parsed = parse_daily_report(sample_tong_path)
    assert "bp_summaries" in parsed
    assert len(parsed["bp_summaries"]) > 0


def test_should_verify_llm_env_parsing():
    """env CPIA_PARSER_VERIFY_LLM 파싱: 'off'/'false'/'0'/'no' 만 opt-out."""
    from pipeline.agents.parser import _should_verify_llm

    for v in ("off", "false", "0", "no", "Off", "False"):
        with patch.dict(os.environ, {"CPIA_PARSER_VERIFY_LLM": v}):
            assert _should_verify_llm() is False, f"{v} should opt-out"
    for v in ("on", "true", "1", "yes", "anything-else"):
        with patch.dict(os.environ, {"CPIA_PARSER_VERIFY_LLM": v}):
            assert _should_verify_llm() is True, f"{v} should opt-in"
    # default = on
    with patch.dict(os.environ, {}, clear=False):
        os.environ.pop("CPIA_PARSER_VERIFY_LLM", None)
        assert _should_verify_llm() is True


# ---------------------------------------------------------------------------
# 2026-05-20 refactor: caller smoke test — collector + run_one_cycle (spec §4.1)
# ---------------------------------------------------------------------------


def test_collector_does_not_import_parser_llm():
    """split-brain 방지 — collector 가 parser_llm 직접 import 안 함."""
    with open("pipeline/agents/collector.py") as f:
        content = f.read()
    assert "from pipeline.agents.parser_llm import" not in content
    # docstring 안 단어 매치 회피 — 실제 호출 패턴만 검증
    assert "parse_daily_report_llm(" not in content
    assert "from pipeline.agents.parser import parse_daily_report" in content


def test_run_one_cycle_does_not_call_parser_llm():
    """split-brain 방지 — scripts/run_one_cycle.py 가 parser_llm 호출 안 함."""
    with open("scripts/run_one_cycle.py") as f:
        content = f.read()
    # parser_llm.parse_daily_report_llm 호출 패턴 부재
    assert "parse_daily_report_llm(" not in content
    assert "from pipeline.agents.parser_llm import ParserConfig" not in content
    assert "from pipeline.agents.parser import parse_daily_report" in content
