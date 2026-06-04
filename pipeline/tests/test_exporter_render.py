"""render_xlsx 단위 — 가짜 SheetData → openpyxl bytes 재 load → cell 검증.

Spec: §7.2, §8.1. LLM 0. PG 0.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from pipeline.agents.exporter import SheetData, render_xlsx


@pytest.fixture
def fake_data() -> SheetData:
    return SheetData(
        output_manpower=[
            {
                "zone": "FAB",
                "disc_abbrev": "TE",
                "bp_code": "PH_X-003",
                "company": "아스타",
                "scope": "PJT",
                "work_text_yesterday": "전일",
                "work_text_today": "금일",
                "work_text_tomorrow": "명일",
                "workers_yesterday_incr": {
                    "morning": 0,
                    "daytime": 1,
                    "overtime": 0,
                    "night": 0,
                    "all_night": 0,
                    "total": 1,
                },
                "workers_yesterday_cum": {
                    "morning": 0,
                    "daytime": 2,
                    "overtime": 0,
                    "night": 0,
                    "all_night": 0,
                    "total": 2,
                },
                "workers_today_incr": {
                    "morning": 0,
                    "daytime": 3,
                    "overtime": 0,
                    "night": 0,
                    "all_night": 0,
                    "total": 3,
                },
                "workers_today_cum": {
                    "morning": 0,
                    "daytime": 4,
                    "overtime": 0,
                    "night": 0,
                    "all_night": 0,
                    "total": 4,
                },
            }
        ],
        occupation_manpower=[
            {
                "disc_abbrev": "TE",
                "bp_code": "PH_X-002",
                "job_type": "BP소장",
                "manager_code": "Mm01",
                "worker_code": "sk01",
                "workers_cum_t": 35,
            },
        ],
        equipment=[
            {
                "disc_abbrev": "공통가설",
                "bp_code": "PH_X-002",
                "equipment_big": "굴삭운반",
                "equipment_mid": "Back Hoe",
                "equipment_spec": "2㎥ 미만",
                "count_cum_t": 0,
            },
        ],
        quantity=[
            {
                "zone": "FAB",
                "trade": "TE",
                "bp_code": "PH_X-003",
                "company": "아스타",
                "scope": None,
                "work_content": "M15 환경유지",
                "unit": "인",
                "plan_total_qty": 100,
                "yesterday_cum_qty": 5,
                "today_qty": 3,
            },
        ],
    )


def test_render_xlsx_returns_bytes(fake_data: SheetData) -> None:
    out = render_xlsx(fake_data, date(2025, 9, 26), project="M15X")
    assert isinstance(out, bytes)
    assert len(out) > 1000


def test_render_xlsx_sheets_present(fake_data: SheetData) -> None:
    out = render_xlsx(fake_data, date(2025, 9, 26), project="M15X")
    wb = load_workbook(BytesIO(out))
    assert wb.sheetnames == ["1.출력인원 현황", "2.직종별 인원현황", "3.장비현황", "4.물량"]


def test_render_xlsx_output_manpower_row(fake_data: SheetData) -> None:
    """시트1 r5 에 zone='FAB' 박혀야 함 (COL dict 가 정확한지 검증)."""
    out = render_xlsx(fake_data, date(2025, 9, 26), project="M15X")
    wb = load_workbook(BytesIO(out))
    ws = wb["1.출력인원 현황"]
    found = any(ws.cell(5, c).value == "FAB" for c in range(1, ws.max_column + 1))
    assert found, (
        f"r5 row 어디에도 'FAB' 없음 — COL dict 확인. 실 r5: {[ws.cell(5, c).value for c in range(1, ws.max_column + 1)]}"
    )


def test_render_xlsx_quantity_row(fake_data: SheetData) -> None:
    """시트4 r5 에 work_content 'M15 환경유지' 박혀야 함."""
    out = render_xlsx(fake_data, date(2025, 9, 26), project="M15X")
    wb = load_workbook(BytesIO(out))
    ws = wb["4.물량"]
    found = any(ws.cell(5, c).value == "M15 환경유지" for c in range(1, ws.max_column + 1))
    assert found


def test_render_xlsx_empty_sheet_keeps_header(fake_data: SheetData) -> None:
    """empty 시트 데이터 (e.g. equipment=[]) 도 헤더만 남아 깨지지 X."""
    fake_data.equipment.clear()
    out = render_xlsx(fake_data, date(2025, 9, 26), project="M15X")
    wb = load_workbook(BytesIO(out))
    ws = wb["3.장비현황"]
    assert ws["B1"].value is not None or ws.max_row >= 4


def test_reconstruct_output_manpower_per_company() -> None:
    from pipeline.agents.exporter import _reconstruct_output_manpower

    agg = [
        {"company": "OTIS", "trade": "건축", "workers_total": 27.3, "work_today": "승강로 작업"},
        {"company": "광건", "trade": "건축", "workers_total": 12.0, "work_today": "면 작업"},
    ]
    out = _reconstruct_output_manpower(agg, {"OTIS": "FAB", "광건": "FAB"})
    assert len(out) == 2
    assert out[0]["company"] == "OTIS"
    assert out[0]["bp_code"] is None
    assert out[0]["zone"] == "FAB"
    assert out[0]["disc_abbrev"] == "건축"
    assert out[0]["work_text_today"] == "승강로 작업"
    assert out[0]["work_text_yesterday"] is None
    assert out[0]["workers_today_incr"] == {"total": 27}


def test_render_xlsx_sheet1_bp_is_company() -> None:
    """시트1 col5(BP사) = 업체명(코드 아님)."""
    from pipeline.agents.exporter import SheetData, render_xlsx, COL_OUTPUT_MANPOWER
    from openpyxl import load_workbook
    from io import BytesIO
    from datetime import date

    sd = SheetData(
        output_manpower=[
            {
                "company": "광건",
                "bp_code": "PH_X-070",
                "disc_abbrev": "건축",
                "work_text_today": "x",
                "workers_today_incr": {"total": 5},
            }
        ],
        occupation_manpower=[],
        equipment=[],
        quantity=[],
        company_wa={"광건": "FAB"},
    )
    wb = load_workbook(BytesIO(render_xlsx(sd, date(2025, 9, 24), project="M15X")))
    ws = wb["1.출력인원 현황"]
    assert ws.cell(5, COL_OUTPUT_MANPOWER["bp_code"]).value == "광건"
