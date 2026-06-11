"""통합작업일보 export 양식 정합 단위 — fake SheetData / render_xlsx → openpyxl 재load. LLM 0, PG 0."""

from __future__ import annotations


from pipeline.agents.exporter import (
    SheetData,
    _company_trade_map,
)


def test_company_trade_map_most_common():
    rows = [
        {"company": "광건", "trade": "건축"},
        {"company": "광건", "trade": "건축"},
        {"company": "장평건설", "trade": "토목"},
        {"company": "장평건설", "trade": "토목"},
        {"company": "장평건설", "trade": "건축"},  # 토목 2 > 건축 1
        {"company": None, "trade": "건축"},  # company 없으면 무시
        {"company": "선보", "trade": None},  # trade 없으면 무시
    ]
    m = _company_trade_map(rows)
    assert m["광건"] == "건축"
    assert m["장평건설"] == "토목"  # 최빈(토목 2 > 건축 1)
    assert None not in m
    assert "선보" not in m
    # 동률 결정론: 토목 vs 건축 1:1 → 사전순(건축 < 토목)
    assert (
        _company_trade_map([{"company": "X", "trade": "토목"}, {"company": "X", "trade": "건축"}])["X"]
        == "건축"
    )
    assert _company_trade_map([]) == {}  # 빈 입력 → 빈 맵


def test_sheetdata_company_trade_default_empty():
    sd = SheetData(output_manpower=[], occupation_manpower=[], equipment=[], quantity=[])
    assert sd.company_trade == {}


# ──────────────────────────────────────────────────────────────────────────────
# Task 2 — 구분(col2) / 시트1 공종 company_trade 폴백
# ──────────────────────────────────────────────────────────────────────────────

from datetime import date  # noqa: E402
from io import BytesIO  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from pipeline.agents.exporter import render_xlsx  # noqa: E402


def _sd_one_company():
    return SheetData(
        output_manpower=[
            {
                "company": "광건",
                "bp_code": "PH_X-070",
                "zone": "CUB",
                "disc_abbrev": None,  # 공종 비어있음 → company_trade 폴백
                "work_text_today": "천장 설치",
                "workers_today_incr": {"total": 5},
            }
        ],
        occupation_manpower=[],
        equipment=[],
        quantity=[
            {
                "company": "광건",
                "trade": "건축",
                "mapper_wa": "CUB",
                "work_content": "경량벽체",
                "unit": "m2",
                "today_qty": 9,
            }
        ],
        company_wa={"광건": "CUB"},
        company_trade={"광건": "건축"},
    )


def test_render_division_col2_first_row_only():
    """구분(col2) = project, 첫 데이터 행(r5)만. 시트1·4 모두."""
    from pipeline.agents.exporter import COL_OUTPUT_MANPOWER, COL_QUANTITY

    wb = load_workbook(BytesIO(render_xlsx(_sd_one_company(), date(2025, 9, 25), project="M15X Ph-3 IBL")))
    # rename 후 시트1 = '1.작업사항'
    ws1 = wb["1.작업사항"]
    assert ws1.cell(5, COL_OUTPUT_MANPOWER["project"]).value == "M15X Ph-3 IBL"
    ws4 = wb["4.물량"]
    assert ws4.cell(5, COL_QUANTITY["project"]).value == "M15X Ph-3 IBL"


def test_render_sheet1_trade_fallback_to_company_trade():
    """시트1 공종(col4): disc_abbrev 없으면 company_trade 폴백."""
    from pipeline.agents.exporter import COL_OUTPUT_MANPOWER

    wb = load_workbook(BytesIO(render_xlsx(_sd_one_company(), date(2025, 9, 25), project="M15X Ph-3 IBL")))
    ws1 = wb["1.작업사항"]
    assert ws1.cell(5, COL_OUTPUT_MANPOWER["trade"]).value == "건축"


def test_unhide_and_clean_removes_hidden_outline_filter():
    import openpyxl
    from pipeline.agents.exporter import _unhide_and_clean

    ws = openpyxl.Workbook().active
    ws.auto_filter.ref = "B5:Y100"
    ws.row_dimensions[8].hidden = True
    ws.row_dimensions[8].outlineLevel = 2
    ws.column_dimensions["D"].hidden = True
    ws.freeze_panes = "F26"
    _unhide_and_clean(ws, n=10)
    assert ws.auto_filter.ref is None
    assert ws.row_dimensions[8].hidden is False
    assert ws.row_dimensions[8].outlineLevel == 0
    assert ws.column_dimensions["D"].hidden is False
    assert ws.column_dimensions["D"].outlineLevel == 0
    assert ws.sheet_view.showOutlineSymbols is False
    assert ws.freeze_panes == "A5"


def test_dedup_clears_same_group_keeps_different_zone():
    import openpyxl
    from pipeline.agents.exporter import _dedup_and_border, DATA_START_ROW
    ws = openpyxl.Workbook().active
    # 시트4 형태: col3=zone, col4=공종, col5=bp, col6=업체. r5/r6 같은 그룹, r7 zone 다름
    data = [
        ("CUB", "건축", "PH_X-070", "광건"),
        ("CUB", "건축", "PH_X-070", "광건"),  # 직전과 전부 같음 → 비움
        ("FAB", "건축", "PH_X-070", "광건"),  # zone 다름 → 유지
    ]
    for i, (z, t, b, c) in enumerate(data):
        r = DATA_START_ROW + i
        ws.cell(r, 3, z); ws.cell(r, 4, t); ws.cell(r, 5, b); ws.cell(r, 6, c)
    _dedup_and_border(ws, n=3, c0=2, c1=17, key_cols=(3, 4, 5, 6))
    # r6(둘째) 키 비움
    assert ws.cell(DATA_START_ROW + 1, 3).value is None
    assert ws.cell(DATA_START_ROW + 1, 6).value is None
    # r7 zone 다르므로 유지
    assert ws.cell(DATA_START_ROW + 2, 3).value == "FAB"
    assert ws.cell(DATA_START_ROW + 2, 6).value == "광건"


def test_dedup_border_group_boundary_thin_internal_hair():
    import openpyxl
    from pipeline.agents.exporter import _dedup_and_border, DATA_START_ROW
    ws = openpyxl.Workbook().active
    for i, c in enumerate(["광건", "광건", "남웅"]):
        ws.cell(DATA_START_ROW + i, 6, c); ws.cell(DATA_START_ROW + i, 3, "CUB")
        ws.cell(DATA_START_ROW + i, 4, "건축"); ws.cell(DATA_START_ROW + i, 5, "B")
    _dedup_and_border(ws, n=3, c0=2, c1=17, key_cols=(3, 4, 5, 6))
    assert ws.cell(DATA_START_ROW + 0, 6).border.top.style == "thin"   # 첫 그룹 시작
    assert ws.cell(DATA_START_ROW + 1, 6).border.top.style == "hair"   # 같은 광건 내부
    assert ws.cell(DATA_START_ROW + 2, 6).border.top.style == "thin"   # 남웅 새 그룹
    assert ws.cell(DATA_START_ROW + 0, 6).border.bottom.style == "hair"


def test_dedup_border_skips_merged_cell():
    """border 범위에 병합셀이 있어도 크래시 없이 dedup·테두리 적용(MergedCell read-only 회피)."""
    import openpyxl
    from pipeline.agents.exporter import _dedup_and_border, DATA_START_ROW
    ws = openpyxl.Workbook().active
    for i in range(2):
        ws.cell(DATA_START_ROW + i, 3, "CUB")
        ws.cell(DATA_START_ROW + i, 6, "광건")
    ws.merge_cells(start_row=DATA_START_ROW, start_column=8, end_row=DATA_START_ROW, end_column=9)
    _dedup_and_border(ws, n=2, c0=2, c1=17, key_cols=(3, 4, 5, 6))  # 예외 없어야
    assert ws.cell(DATA_START_ROW + 1, 6).value is None  # 둘째 행 dedup 정상


def test_unmerge_data_region_clears_intersecting_merges():
    import openpyxl
    from pipeline.agents.exporter import _unmerge_data_region, SheetData, DATA_START_ROW
    wb = openpyxl.Workbook()
    wb.active.title = "1.출력인원 현황"
    ws2 = wb.create_sheet("2.직종별 인원현황")
    wb.create_sheet("3.장비현황")
    wb.create_sheet("4.물량")
    # occupation 5행(r5~r9) 가정: r6 데이터구역 안 병합 → 해제, r50 밖 병합 → 유지
    ws2.merge_cells(start_row=DATA_START_ROW + 1, start_column=2, end_row=DATA_START_ROW + 1, end_column=6)
    ws2.merge_cells(start_row=50, start_column=2, end_row=50, end_column=6)
    ws2.merge_cells(start_row=4, start_column=2, end_row=DATA_START_ROW, end_column=6)  # 헤더(r4)에서 시작 → 보존돼야
    sd = SheetData(output_manpower=[], occupation_manpower=[{} for _ in range(5)], equipment=[], quantity=[])
    _unmerge_data_region(wb, sd)
    refs = {str(m) for m in ws2.merged_cells.ranges}
    assert "B6:F6" not in refs       # 데이터구역(r6) 병합 해제됨
    assert "B50:F50" in refs          # 데이터구역 밖(r50) 유지
    assert "B4:F5" in refs            # 헤더에서 시작한 병합은 보존(I-1)


def test_finalize_sheet1_renamed_and_headcount_hidden():
    wb = load_workbook(BytesIO(render_xlsx(_sd_one_company(), date(2025, 9, 25), project="M15X Ph-3 IBL")))
    assert wb.sheetnames == ["1.작업사항", "2.직종별 인원현황", "3.장비현황", "4.물량"]
    ws1 = wb["1.작업사항"]
    assert ws1.column_dimensions["J"].hidden is True   # col10 인원
    assert ws1.column_dimensions["AK"].hidden is True  # col37
    assert ws1.column_dimensions["I"].hidden is False  # col9 명일작업(주요작업까지 보임)


def test_finalize_sheet2_no_hidden_no_filter_white():
    sd = _sd_one_company()
    sd.occupation_manpower = [
        {"disc_abbrev": "건축", "company": "광건", "job_type": "비계공", "workers_cum_t": 3},
        {"disc_abbrev": "건축", "company": "광건", "job_type": "형틀목공", "workers_cum_t": 5},
    ]
    wb = load_workbook(BytesIO(render_xlsx(sd, date(2025, 9, 25), project="M15X Ph-3 IBL")))
    ws2 = wb["2.직종별 인원현황"]
    assert ws2.auto_filter.ref is None
    assert ws2.row_dimensions[5].hidden is False
    assert ws2.row_dimensions[6].hidden is False
    assert ws2.freeze_panes == "A6"  # 시트2 헤더 3행(r3+r4+r5 서브헤더) → 데이터 r6
    assert ws2.column_dimensions["D"].hidden is False  # 템플릿 숨김열 D 풀림
    assert ws2.column_dimensions["Y"].hidden is False  # 템플릿 숨김열 Y 풀림
    # r5 는 이제 서브헤더 행 — _write_subheaders 가 r4 헤더 스타일(fill 포함)을 복제하므로
    # 회색 아티팩트 방지 단언 제거(헤더 fill 이 전달됨은 의도된 동작).


def test_finalize_empty_sheet_safe():
    """빈 시트(occupation/equipment=[]) 도 _finalize 가 안 깨짐."""
    sd = _sd_one_company()
    sd.occupation_manpower = []
    sd.equipment = []
    wb = load_workbook(BytesIO(render_xlsx(sd, date(2025, 9, 25), project="M15X Ph-3 IBL")))
    assert "2.직종별 인원현황" in wb.sheetnames


def test_sheet1_pjt_to_bp_top_aligned():
    """시트1 작업사항: PJT(col2)~BP사(col5) 세로 위 맞춤 (템플릿 기본 아래맞춤 교정)."""
    wb = load_workbook(BytesIO(render_xlsx(_sd_one_company(), date(2025, 9, 25), project="M15X Ph-3 IBL")))
    ws1 = wb["1.작업사항"]
    for c in (2, 3, 4, 5):
        assert ws1.cell(5, c).alignment.vertical == "top", f"col{c} vertical={ws1.cell(5, c).alignment.vertical}"


# ── 2026-06-11 시트2·3 r5 서브헤더 복원 + 유령열(Y/25) 제거 ──
def _sd_two_occ_eq():
    """시트2·3 데이터 2행씩 — r6 시작·서브헤더·유령열 검증용."""
    sd = _sd_one_company()
    sd.occupation_manpower = [
        {"disc_abbrev": "건축", "company": "광건", "job_type": "비계공", "workers_cum_t": 3},
        {"disc_abbrev": "건축", "company": "광건", "job_type": "형틀목공", "workers_cum_t": 5},
    ]
    sd.equipment = [
        {"disc_abbrev": "건축", "company": "광건", "equipment_big": "굴삭운반기계류",
         "equipment_mid": "Back Hoe", "equipment_spec": "2㎥ 미만", "count_cum_t": 1},
    ]
    return sd


def _rendered_two():
    return load_workbook(BytesIO(render_xlsx(_sd_two_occ_eq(), date(2025, 9, 25), project="M15X Ph-3 IBL")))


def test_sheet23_r5_subheader_restored():
    """시트2·3 r5 = 좌측라벨 + 조출/주간/연장/야간/철야/누계(×3). 최종(Y/25) 없음."""
    wb = _rendered_two()
    ws2 = wb["2.직종별 인원현황"]
    assert ws2.cell(5, 2).value == "Dis"
    assert ws2.cell(5, 3).value == "협력업체"
    assert ws2.cell(5, 8).value == "주간"      # H5
    assert ws2.cell(5, 10).value == "야간"     # J5
    assert ws2.cell(5, 12).value == "Cum. T"   # L5 전일누계 합
    assert ws2.cell(5, 18).value == "Incr. T"  # R5 금일투입 합
    assert ws2.cell(5, 24).value == "Cum. T"   # X5 금일누계 합
    assert ws2.cell(5, 25).value is None        # Y5 최종 = 미작성
    ws3 = wb["3.장비현황"]
    assert ws3.cell(5, 4).value == "대분류"     # 시트3 좌측 라벨 차이
    assert ws3.cell(5, 8).value == "주간"


def test_sheet23_data_starts_row6():
    """시트2·3 데이터는 r6 부터 (r5 는 서브헤더)."""
    wb = _rendered_two()
    ws2 = wb["2.직종별 인원현황"]
    assert ws2.cell(6, 2).value == "건축"       # 첫 데이터 행 = 공종(disc_abbrev)
    assert ws2.cell(6, 3).value == "광건"
    assert ws2.cell(6, 2).fill.patternType is None  # 데이터행 흰색(회색 아티팩트 방지 — 구 r5 가드 이전)
    ws3 = wb["3.장비현황"]
    assert ws3.cell(6, 4).value == "굴삭운반기계류"  # 첫 데이터 = 대분류


def test_sheet23_no_ghost_column_25():
    """유령열 제거: 시트2·3 Y열(25) 데이터행에 값·테두리 모두 없음."""
    from openpyxl.cell.cell import MergedCell
    wb = _rendered_two()
    for sname, ndata in (("2.직종별 인원현황", 2), ("3.장비현황", 1)):
        ws = wb[sname]
        for i in range(ndata):
            cell = ws.cell(6 + i, 25)
            assert cell.value is None
            if not isinstance(cell, MergedCell):
                b = cell.border
                assert not any(s and s.style for s in (b.left, b.right, b.top, b.bottom)), \
                    f"{sname} Y{6+i} 에 테두리 잔존(유령열)"


def test_sheet14_data_still_row5_no_regression():
    """시트1·4 는 종전대로 r5 시작 (회귀 방지)."""
    from pipeline.agents.exporter import COL_OUTPUT_MANPOWER, COL_QUANTITY
    wb = _rendered_two()
    assert wb["1.작업사항"].cell(5, COL_OUTPUT_MANPOWER["project"]).value == "M15X Ph-3 IBL"
    assert wb["4.물량"].cell(5, COL_QUANTITY["project"]).value == "M15X Ph-3 IBL"


def test_sheet4_pct_column_no_red_conditional_format():
    """시트4 F%(N열) '>1 → 빨강' 조건부서식 제거 — F% 는 퍼센트라 전 셀 빨강 방지(흰색)."""
    wb = _rendered_two()
    assert len(list(wb["4.물량"].conditional_formatting)) == 0
    # 타 시트 템플릿 조건부서식은 보존(시트1 manpower 그리드 CF 무손상).
    assert len(list(wb["1.작업사항"].conditional_formatting)) > 0
