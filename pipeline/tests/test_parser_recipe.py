from __future__ import annotations

import math
import openpyxl
import pytest
from openpyxl import load_workbook
from pipeline.agents.parser import ParseError, parse_daily_report
from pipeline.agents.recipe.bpcode import is_bp_code_v2, looks_like_code_no_korean
from pipeline.agents.recipe.extract import (
    extract_sheet1,
    extract_sheet2,
    extract_sheet3,
    extract_sheet4,
)
from pipeline.agents.recipe import parse_with_recipe
from pipeline.agents.recipe.fingerprint import fingerprint
from pipeline.agents.recipe.model import FormatRecipe, SheetRecipe, REQUIRED_FIELDS
from pipeline.agents.recipe.recognize import load_cached_recipe, recognize_format, save_recipe
from pipeline.agents.recipe.synonyms import field_for_header, role_for_sheet
from scripts.synth_formats import make_variant

UNI = "docs/sample/통합작업일보_M15X Ph-3 IBL PJT 작업일보_20250925.xlsx"
BP = "docs/sample/업체작업일보_M15X_PH4_041_M15X_PH4_UT_ 중원엔지니어링_UT배관공사_FAB_250917.xlsx"


def test_recipe_model_construct():
    sr = SheetRecipe(
        sheet_name="4.물량",
        columns={"work_content": 7, "unit": 8, "today_qty": 11, "trade": 4, "bp": 5},
        timeshift_blocks={},
        group_boundary="bp_company_pair",
        header_rows=(3, 4),
        final_flag_col=None,
    )
    fr = FormatRecipe(
        fingerprint="abc",
        format_class="업체",
        sheets={"activity": sr},
        generated_by="deterministic",
        confidence=1.0,
    )
    assert fr.sheets["activity"].columns["today_qty"] == 11
    assert "today_qty" in REQUIRED_FIELDS["activity"]


def test_fingerprint_distinguishes_and_is_stable():
    wb_u = openpyxl.load_workbook(UNI, data_only=True)
    wb_b = openpyxl.load_workbook(BP, data_only=True)
    fu1 = fingerprint(wb_u)
    fu2 = fingerprint(wb_u)
    fb = fingerprint(wb_b)
    assert fu1 == fu2  # 같은 wb 두 번 → 동일 (결정성)
    assert fu1 != fb  # 통합 vs 업체 → 다름
    assert len(fu1) == 64  # sha256 hex
    wb_u.close()
    wb_b.close()


def test_synonyms_map_headers_and_sheets():
    assert field_for_header("금일실적") == "today_qty"
    assert field_for_header("당일실적") == "today_qty"  # V2 동의어
    assert field_for_header("대표물량_작업내용") == "work_content"
    assert field_for_header("작업내용") == "work_content"
    assert field_for_header("BP사") == "bp"
    assert field_for_header("업체명") == "company_col"
    # sheet2/3 필수필드 (실측 헤더)
    assert field_for_header("관리자/근로자") == "disc_code"  # sheet2 col4 (Mm01 앵커)
    assert field_for_header("직종") == "job_type"  # sheet2 col6
    assert field_for_header("대분류") == "equip_big"  # sheet3 col4
    assert field_for_header("중분류") == "equip_mid"
    assert field_for_header("규 격") == "equip_spec"  # _norm_header 가 공백 제거 → "규격"
    # sheet1 작업텍스트 (헤더에 날짜 brackets 붙음 → substring)
    assert field_for_header("전일 작업사항 [09/24]") == "work_prev"
    assert field_for_header("금일 작업사항 [09/25]") == "work_today"
    assert field_for_header("명일 작업사항 [09/26]") == "work_tomorrow"
    assert field_for_header("ZZZ알수없음") is None
    assert role_for_sheet("4.물량") == "activity"
    assert role_for_sheet("물량집계") == "activity"  # V1 시트명 변이 alias
    assert role_for_sheet("1.출력인원 현황") == "bp_summary"
    assert role_for_sheet("2.직종별 인원현황") == "job"
    assert role_for_sheet("3.장비현황") == "equip"


def test_recognize_integrated_format():
    wb = openpyxl.load_workbook(UNI, data_only=True)
    r = recognize_format(wb)
    assert r.format_class == "통합"
    assert r.generated_by == "deterministic"
    assert r.confidence == 1.0
    # sheet4 컬럼 (통합 실측: company_col=6 이 있어 이후 1칸 우측 시프트)
    # trade=4 bp=5 company_col=6 scope=7 work_content=8 unit=9 today_qty=12
    a = r.sheets["activity"]
    assert a.columns["trade"] == 4 and a.columns["bp"] == 5
    assert a.columns["company_col"] == 6 and a.columns["work_content"] == 8
    assert a.columns["unit"] == 9 and a.columns["today_qty"] == 12
    # sheet2 final_flag 존재(통합)
    assert r.sheets["job"].final_flag_col == 25
    # sheet2 timeshift (실측: today_cum block 첫칸 col19)
    assert r.sheets["job"].timeshift_blocks["today_cum"] == 19
    # sheet3 final_flag 도 통합은 'O' 게이트 존재 (실측 col25)
    assert r.sheets["equip"].final_flag_col == 25
    wb.close()


def test_recognize_vendor_format():
    wb = openpyxl.load_workbook(BP, data_only=True)
    r = recognize_format(wb)
    assert r.format_class == "업체"
    a = r.sheets["activity"]
    assert "company_col" not in a.columns  # 업체는 업체명 컬럼 없음
    assert a.columns["work_content"] == 7 and a.columns["today_qty"] == 11
    assert r.sheets["job"].final_flag_col is None  # 업체는 'O' 게이트 없음
    wb.close()


def test_bp_code_v2_accepts_variants():
    # 기존 통합/업체 (회귀)
    assert is_bp_code_v2("PH_X-088")
    assert is_bp_code_v2("M15X_PH4_041")
    # P1 변종 — 누출 차단 대상
    assert is_bp_code_v2("M15X-PH4-041")  # dash
    assert is_bp_code_v2("M15X_PH4_041A")  # trailing suffix
    # 회사명은 코드 아님
    assert not is_bp_code_v2("중원엔지니어링")
    assert not is_bp_code_v2("협력업체01")


def test_leak_guard_flags_ascii_only_codeish():
    # 한글 0글자 + 코드 모양 → 회사명 아님 (격리 대상)
    assert looks_like_code_no_korean("ABC-2025-001")
    assert not looks_like_code_no_korean("중원엔지니어링")
    assert not looks_like_code_no_korean("Hyundai E&C")  # 공백 포함 영문 회사명 — 코드 모양 아님


def _activities_via_recipe(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    r = recognize_format(wb)
    acts = extract_sheet4(wb[r.sheets["activity"].sheet_name], r.sheets["activity"])
    wb.close()
    return acts


def test_sheet4_equiv_integrated():
    golden = parse_daily_report(UNI)["activities"]
    got = _activities_via_recipe(UNI)
    assert len(got) == len(golden) == 1193
    assert got == golden  # dict 완전 일치


def test_sheet4_equiv_vendor():
    golden = parse_daily_report(BP)["activities"]
    got = _activities_via_recipe(BP)
    assert len(got) == len(golden) == 33
    assert got == golden


def _bp_via_recipe(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    r = recognize_format(wb)
    out = extract_sheet1(wb[r.sheets["bp_summary"].sheet_name], r.sheets["bp_summary"])
    wb.close()
    return out


def test_sheet1_equiv_integrated():
    golden = parse_daily_report(UNI)["bp_summaries"]
    got = _bp_via_recipe(UNI)
    assert len(got) == len(golden) == 99
    assert got == golden


def test_sheet1_equiv_vendor_company():
    got = _bp_via_recipe(BP)
    assert {s["company"] for s in got} == {"중원엔지니어링"}  # 누출 0, 정규화 OK


def test_sheet2_equiv():
    for path, n in [(UNI, 975), (BP, 76)]:
        golden = parse_daily_report(path)["job_allocations"]
        wb = openpyxl.load_workbook(path, data_only=True)
        r = recognize_format(wb)
        got = extract_sheet2(wb[r.sheets["job"].sheet_name], r.sheets["job"])
        wb.close()
        assert len(got) == len(golden) == n
        assert got == golden


def test_sheet3_equiv():
    for path, n in [(UNI, 1111), (BP, 68)]:
        golden = parse_daily_report(path)["equipment_allocations"]
        wb = openpyxl.load_workbook(path, data_only=True)
        r = recognize_format(wb)
        got = extract_sheet3(wb[r.sheets["equip"].sheet_name], r.sheets["equip"])
        wb.close()
        assert len(got) == len(golden) == n
        assert got == golden


def test_full_report_equiv_both_formats():
    for path in (UNI, BP):
        golden = parse_daily_report(path)
        got = parse_with_recipe(path)
        for key in (
            "activities",
            "bp_summaries",
            "job_allocations",
            "equipment_allocations",
            "report_date",
            "project",
            "worker_outputs",
            "report_meta",
        ):
            assert got[key] == golden[key], f"{path} {key} 불일치"


def test_company_mismatch_sets_review():
    demo = "data/demo_inputs/20250925/업체작업일보_M15X_PH3_남웅건설_단열공사_20250925.xlsm"
    got = parse_with_recipe(demo)
    assert {a["company"] for a in got["activities"]} == {"중원엔지니어링"}  # 셀 SSOT, 덮어쓰기 X
    assert all(a["requires_review"] for a in got["activities"])  # 셀 company 가 파일명에 없음 → 전부 review


def test_recipe_cache_roundtrip(tmp_path, monkeypatch):
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    wb = openpyxl.load_workbook(BP, data_only=True)
    r = recognize_format(wb)
    wb.close()
    save_recipe(r)
    loaded = load_cached_recipe(r.fingerprint)
    assert loaded is not None
    assert loaded.format_class == r.format_class
    assert loaded.sheets["activity"].columns == r.sheets["activity"].columns
    assert loaded.sheets["job"].header_rows == r.sheets["job"].header_rows  # tuple 복원
    assert isinstance(loaded.sheets["job"].header_rows, tuple)


class _MockLLM:
    def __init__(self, payload):
        self.payload = payload
        self.calls = 0

    def complete(self, prompt: str) -> str:
        self.calls += 1
        return self.payload


def _make_activity_gap_wb(tmp_path):
    # 업체 파일의 activity 시트 work_content 헤더를 미인식 문자열로 바꿔 Tier1 gap 유발
    wb = load_workbook(BP)
    # recognize 로 activity 시트명/헤더행/work_content 컬럼을 실측해 그 헤더 셀을 깨뜨림
    probe = openpyxl.load_workbook(BP, data_only=True)
    r0 = recognize_format(probe, use_cache=False)
    a = r0.sheets["activity"]
    wc_col = a.columns["work_content"]
    ws = wb[a.sheet_name]
    # 헤더행(3,4) 중 work_content 헤더가 있는 행을 gibberish 로 치환
    for hr in a.header_rows:
        ws.cell(hr, wc_col).value = "ZZZUNKNOWN"
    probe.close()
    out = tmp_path / "gap.xlsx"
    wb.save(out)
    return str(out), wc_col


def test_tier2_broken_header_yields_tier1_gap(tmp_path, monkeypatch):
    # llm 없이 Tier1 만: work_content gap → deterministic + confidence<1.0 (fallback 가 진짜 발동함을 증명)
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    path, wc_col = _make_activity_gap_wb(tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    r = recognize_format(wb, use_cache=False, strict=False)
    wb.close()
    assert r.generated_by == "deterministic"
    assert r.confidence < 1.0
    assert "work_content" not in r.sheets["activity"].columns


def test_tier2_llm_fills_gap_deterministic(tmp_path, monkeypatch):
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    path, wc_col = _make_activity_gap_wb(tmp_path)
    payload = '{"activity": {"work_content": %d}}' % wc_col

    wb1 = openpyxl.load_workbook(path, data_only=True)
    m1 = _MockLLM(payload)
    r1 = recognize_format(wb1, use_cache=False, strict=False, llm_client=m1)
    wb1.close()
    assert m1.calls == 1
    assert r1.generated_by == "llm"
    assert r1.sheets["activity"].columns["work_content"] == wc_col

    # 결정성: 같은 입력 2회 → 동일 컬럼맵
    wb2 = openpyxl.load_workbook(path, data_only=True)
    m2 = _MockLLM(payload)
    r2 = recognize_format(wb2, use_cache=False, strict=False, llm_client=m2)
    wb2.close()
    assert r1.sheets["activity"].columns == r2.sheets["activity"].columns


def test_tier2_bad_json_no_cache_poisoning(tmp_path, monkeypatch):
    # LLM 이 JSON 파싱 실패 응답 → gap 그대로. generated_by 는 llm 아님, 캐시 파일도 안 생김
    # (gappy recipe 를 llm 으로 캐시하면 다음 실행이 Tier3 안전망을 영구 우회 = 조용한 오염).
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    path, _ = _make_activity_gap_wb(tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    r = recognize_format(wb, use_cache=False, strict=False, llm_client=_MockLLM("NOT JSON{{"))
    wb.close()
    assert r.generated_by == "deterministic"  # 못 채웠으므로 llm 표기 안 함
    assert "work_content" not in r.sheets["activity"].columns
    assert not list(tmp_path.glob("*.json"))  # gappy recipe 캐시 금지


def test_tier2_non_numeric_value_no_crash(tmp_path, monkeypatch):
    # LLM 이 비숫자 컬럼값 반환 → caller 까지 crash 금지. 해당 필드 skip (Tier3 가 잔여 처리).
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    path, _ = _make_activity_gap_wb(tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    r = recognize_format(
        wb, use_cache=False, strict=False, llm_client=_MockLLM('{"activity": {"work_content": "oops"}}')
    )
    wb.close()
    assert "work_content" not in r.sheets["activity"].columns  # 비숫자 → skip


def test_tier3_loud_fail_on_unmappable_required(tmp_path, monkeypatch):
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    path, _ = _make_activity_gap_wb(tmp_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    with pytest.raises(ParseError) as e:
        recognize_format(wb, llm_client=None)  # strict=True default, no LLM → gap unresolved
    wb.close()
    assert "work_content" in str(e.value)
    # 부분 recipe 가 _review 로 방출됐는지 (정상 캐시는 오염 안 됨)
    review_files = list((tmp_path / "_review").glob("*.json"))
    assert len(review_files) == 1
    assert not list(tmp_path.glob("*.json"))  # 정상 캐시 디렉토리엔 없음


def test_synth_variants_load(tmp_path):
    for v in ("V1", "V2", "V3", "V4", "V5"):
        out = tmp_path / f"{v}.xlsx"
        make_variant(BP, v, str(out))
        wb = openpyxl.load_workbook(out, data_only=True)
        assert len(wb.sheetnames) >= 4
        wb.close()


def _approx_equal(a, b) -> bool:
    """float 필드는 isclose, 나머지는 exact 로 재귀 비교.

    합성 변이는 openpyxl 재저장 과정에서 수식 캐시값(progress_percent 등)이 ~1e-17 수준
    float drift 를 겪는다 (Excel↔openpyxl 직렬화 차이 — 파서 로직과 무관). 의미 등가 검증은
    이 drift 를 isclose 로 흡수해야 정직하다. 결정성 검증(같은 파일 2회)은 exact == 로 한다.
    """
    if isinstance(a, float) and isinstance(b, float):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(_approx_equal(a[k], b[k]) for k in a)
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
        return len(a) == len(b) and all(_approx_equal(x, y) for x, y in zip(a, b))
    return a == b


_SYNTH_ENTITY_KEYS = ("activities", "bp_summaries", "job_allocations", "equipment_allocations")


def test_synth_equivalence_and_determinism(tmp_path, monkeypatch):
    # V1(시트rename)/V2(헤더동의어)/V3(컬럼shuffle) — 라벨·위치만 바뀐 의미등가 변이.
    # parse_with_recipe 산출이 canonical 업체 산출과 의미등가(float 필드 isclose) + 2회 결정성.
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    canon = parse_with_recipe(BP)
    for v in ("V1", "V2", "V3"):
        # 파일명에 회사명(중원엔지니어링)+날짜(250917) → company match(리뷰X) + _extract_meta 날짜 OK
        out = tmp_path / f"업체작업일보_M15X_PH4_중원엔지니어링_{v}_250917.xlsx"
        make_variant(BP, v, str(out))
        got1 = parse_with_recipe(str(out))
        got2 = parse_with_recipe(str(out))
        for key in _SYNTH_ENTITY_KEYS:
            assert _approx_equal(got1[key], canon[key]), f"{v} {key} 의미등가 실패"
            assert got1[key] == got2[key], f"{v} {key} 결정성 실패"


def test_synth_v4_format_flip(tmp_path, monkeypatch):
    # V4: 업체파일에 업체명 컬럼 삽입 → 통합축. recipe 가 format_class='통합' 으로 flip 감지하고
    # 4 엔티티 건수를 보존. 단 activity dict 는 vendor↔통합 추출 경로가 본질적으로 달라(통합 ffill
    # vs vendor group-scan) byte/의미 등가 비대상 — 경로 차이는 파서 결함이 아니라 포맷축 차이다.
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    canon = parse_with_recipe(BP)
    out = tmp_path / "업체작업일보_M15X_PH4_중원엔지니어링_V4_250917.xlsx"
    make_variant(BP, "V4", str(out))
    wb = openpyxl.load_workbook(str(out), data_only=True)
    r = recognize_format(wb, use_cache=False, strict=False)
    wb.close()
    assert r.format_class == "통합"  # 업체명 컬럼 삽입 → 통합 인식 (flip)
    got = parse_with_recipe(str(out))
    for key in _SYNTH_ENTITY_KEYS:
        assert len(got[key]) == len(canon[key]), f"V4 {key} 건수 불일치"


def test_capture_exactness_vs_golden():
    """recipe 경로 집계 == golden(parse_daily_report) 경로 집계.

    - JobAllocation.workers_today_incr["total"]: TimeShiftCounts.total (금일 투입 합계)
    - BPSummary.today_incr_total: 편의 int 필드 (시트1 금일 투입 합계)
    - EquipmentAllocation.count_today_cum["total"]: TimeShiftCounts.total (금일 투입누계 합계)

    주의: Σ(job) != Σ(bp) 는 의도된 차이 (_is_admin_job 제외 = 21 차이). 여기서 검증하는 것은
    recipe 경로 vs golden 경로 간 동일성이지, 시트 간 교차 동일성이 아니다.
    """
    for path in (UNI, BP):
        rep = parse_with_recipe(path)
        gold = parse_daily_report(path)
        assert sum(j["workers_today_incr"]["total"] for j in rep["job_allocations"]) == sum(
            j["workers_today_incr"]["total"] for j in gold["job_allocations"]
        ), f"{path}: job workers_today_incr total 불일치"
        assert sum(b["today_incr_total"] for b in rep["bp_summaries"]) == sum(
            b["today_incr_total"] for b in gold["bp_summaries"]
        ), f"{path}: bp today_incr_total 불일치"
        assert sum(e["count_today_cum"]["total"] for e in rep["equipment_allocations"]) == sum(
            e["count_today_cum"]["total"] for e in gold["equipment_allocations"]
        ), f"{path}: equipment count_today_cum total 불일치"


def test_cache_schema_drift_returns_none(tmp_path, monkeypatch):
    # valid JSON 이지만 스키마 드리프트(키 누락/리네임/top-level non-dict) 캐시 → crash 아니라 None.
    # docstring "손상 시 None" 계약 + public 경로(use_cache=True) brick 방지. (최종리뷰 MEDIUM #1)
    monkeypatch.setattr("pipeline.agents.recipe.recognize.RECIPE_CACHE_DIR", tmp_path)
    (tmp_path / "miss_sheets.json").write_text('{"format_class": "업체"}', encoding="utf-8")
    assert load_cached_recipe("miss_sheets") is None  # KeyError 'sheets' → None
    (tmp_path / "toplevel_list.json").write_text("[1, 2, 3]", encoding="utf-8")
    assert load_cached_recipe("toplevel_list") is None  # TypeError → None
    # public 경로: 실제 fingerprint 자리에 드리프트 파일을 둬도 Tier1 재해석으로 복구(brick X)
    wb = openpyxl.load_workbook(BP, data_only=True)
    fpv = fingerprint(wb)
    (tmp_path / f"{fpv}.json").write_text('{"format_class": "업체"}', encoding="utf-8")
    r = recognize_format(wb)  # use_cache=True default — crash 나면 안 됨
    wb.close()
    assert r.format_class == "업체" and r.confidence == 1.0  # 재해석 복구


def test_synth_v5_english_headers_loud_fail(tmp_path):
    # V5(영문 헤더)는 결정론 흡수 대상이 아님 — _resolve_sheet4_columns 가 한글 토큰만 인식하므로
    # parse_with_recipe(V5) 는 ParseError 로 loud-fail (조용한 오염 X). 영문 변이는 Tier2 LLM 영역.
    # (최종리뷰 MEDIUM #3: V5 한계를 명시 테스트로 박제)
    out = tmp_path / "업체작업일보_M15X_PH4_중원엔지니어링_V5_250917.xlsx"
    make_variant(BP, "V5", str(out))
    with pytest.raises(ParseError) as e:
        parse_with_recipe(str(out))
    assert "work_content" in str(e.value) or "today_qty" in str(e.value)
